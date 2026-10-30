"""Build the editable, evidence-linked three-minute lab02 deck and review assets.

Does not overwrite the original PPTX, report, or earlier recording.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import defaultdict
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_TICK_MARK, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = ROOT
TEST_COUNT = 200
WORK = ROOT / "outputs/qa/lab02-ppt-video-3min"
DELIVERY = ROOT / "deliverables/lab02"
PPTX = DELIVERY / "第31组_推理实验答辩_修订版.pptx"
SUPPLEMENTARY = DELIVERY / "第31组_推理实验答辩_补充材料.pptx"
ORIGINAL = DELIVERY / "第31组_推理实验答辩.pptx"
BG, INK, MUTED = "F8F5F0", "292724", "71685F"
ACCENT, BLUE, LINE, GRAY, PALE = "9A4728", "426D80", "D9CFC3", "939DA0", "EFE6DB"
FONT = "Microsoft YaHei"


def read_csv(relative):
    with (DATA_ROOT / relative).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def text(slide, x, y, width, height, value, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(str(value).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(5)
        paragraph.line_spacing = 1.10
        run = paragraph.add_run()
        run.text = line
        run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)
        properties = run._r.get_or_add_rPr()
        properties.set("lang", "zh-CN")
        east_asian = OxmlElement("a:ea")
        east_asian.set("typeface", FONT)
        properties.append(east_asian)
    return shape


def box(slide, x, y, width, height, fill=PALE, oval=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if oval else MSO_SHAPE.RECTANGLE,
                                  Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, arrow=False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.width = Pt(1.3)
    shape.line.color.rgb = RGBColor.from_string(color)
    if arrow:
        endpoint = OxmlElement("a:tailEnd")
        endpoint.set("type", "triangle")
        shape._element.spPr.get_or_add_ln().append(endpoint)
    return shape


def page(presentation, section, title, source=""):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    text(slide, .65, .61, 12, .24, section, 11, ACCENT, True)
    text(slide, .65, 1.02, 12.03, .64, title, 29, INK, True)
    # Evidence paths stay in the build inputs and QA records, off the slides.
    return slide


def table(slide, x, y, widths, rows, row_height=.57, size=17):
    shape = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y),
                                  Inches(sum(widths)), Inches(row_height * len(rows)))
    result = shape.table
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    for row_index, values in enumerate(rows):
        result.rows[row_index].height = Inches(row_height)
        for column_index, value in enumerate(values):
            cell = result.cell(row_index, column_index)
            cell.text = str(value)
            cell.margin_left = Inches(.12)
            cell.margin_right = Inches(.08)
            cell.margin_top = cell.margin_bottom = Inches(.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(PALE if row_index == 0 else BG)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.04
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(size)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor.from_string(INK if row_index == 0 else MUTED)
                    east_asian = OxmlElement("a:ea")
                    east_asian.set("typeface", FONT)
                    run._r.get_or_add_rPr().append(east_asian)
    return shape


def chart_style(chart, labels=True):
    chart.has_legend = False
    chart.chart_style = 10
    chart.font.name, chart.font.size = FONT, Pt(13)
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.size = Pt(12)
        axis.format.line.color.rgb = RGBColor.from_string(LINE)
        axis.major_tick_mark = XL_TICK_MARK.NONE
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(LINE)
    plot = chart.plots[0]
    plot.has_data_labels = labels
    if labels:
        plot.data_labels.font.name = FONT
        plot.data_labels.font.size = Pt(14)
        plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    for series_index, series in enumerate(chart.series):
        color = [BLUE, ACCENT, GRAY][series_index % 3]
        series.format.line.color.rgb = RGBColor.from_string(color)
        series.format.line.width = Pt(2.5)
        series.format.line.fill.solid()
        series.format.line.fill.fore_color.rgb = RGBColor.from_string(color)
        if chart.chart_type in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS):
            series.marker.style = XL_MARKER_STYLE.CIRCLE if series_index == 0 else XL_MARKER_STYLE.SQUARE
            series.marker.size = 6
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
            series.marker.format.line.color.rgb = RGBColor.from_string(color)
        else:
            for point in series.points:
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(color)
                point.format.line.fill.background()


def add_chart(slide, kind, x, y, width, height, categories, series, labels=True):
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    chart = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(width), Inches(height), data).chart
    chart_style(chart, labels)
    return chart


def lookup(rows, **conditions):
    return next(row for row in rows if all(row[key] == value for key, value in conditions.items()))


def build():
    WORK.mkdir(parents=True, exist_ok=True)
    DELIVERY.mkdir(parents=True, exist_ok=True)
    original_hash = hashlib.sha256(ORIGINAL.read_bytes()).hexdigest()
    narration = json.loads((WORK / "narration.json").read_text(encoding="utf-8"))
    brief = json.loads((ROOT / "outputs/qa/lab02-materials/brief.json").read_text(encoding="utf-8"))
    forward_path = "outputs/lab02/benchmarks/forward_summary.csv"
    score_path = "outputs/lab02/experiments/summary.csv"
    graph_path = "outputs/lab02/end_to_end/summary.csv"
    sensitivity_path = "outputs/lab02/experiments/sensitivity.csv"
    ablation_path = "outputs/lab02/experiments/ablations.csv"
    forward, scores, graph = map(read_csv, [forward_path, score_path, graph_path])
    sensitivity, ablations = map(read_csv, [sensitivity_path, ablation_path])
    benchmark = [lookup(forward, groups="16", model=model) for model in ["naive_scan", "indexed_naive", "semi_naive"]]
    speedup = float(benchmark[1]["elapsed_ms_mean"]) / float(benchmark[2]["elapsed_ms_mean"])
    bayes = lookup(scores, scenario="clean", model="bayes_learned")
    bayes_manual = lookup(scores, scenario="clean", model="bayes_manual")
    fuzzy = lookup(scores, scenario="clean", model="fuzzy_learned")
    graph_rows = [lookup(graph, scenario="false_loop_020", model=model) for model in
                  ["hybrid_adaptive_unlocked", "hybrid_adaptive_locked", "hybrid_adaptive_calibrated"]]
    false_positive_drop = (float(graph_rows[1]["false_positive_rate_mean"]) - float(graph_rows[2]["false_positive_rate_mean"])) * 100
    presentation = Presentation()
    presentation.core_properties.title = "面向企业股权风险评估的确定性与不确定性混合推理"
    presentation.core_properties.subject = "人工智能导论实验二：算法实现、优化与实验验证"
    presentation.core_properties.author = "第31组"
    presentation.core_properties.keywords = "企业股权,知识建模,确定性推理,不确定性推理,混合推理"
    presentation.slide_width, presentation.slide_height = Inches(13.333333), Inches(7.5)

    slide = page(presentation, "人工智能导论 · 实验二", "企业股权风险推理")
    text(slide, .67, 1.90, 11.8, 1.45, "确定性与不确定性混合推理\n算法实现、优化与实验验证", 34, INK, True)
    for xpos, label, detail in [(.70, "确定性", "一阶知识库与完整证明"), (4.93, "不确定性", "Bayes 与 Mamdani"), (9.14, "混合创新", "评分与复核策略分离")]:
        line(slide, xpos, 3.49, xpos + 3.48, 3.49, ACCENT)
        text(slide, xpos, 3.81, 3.55, .48, label, 25, ACCENT, True)
        text(slide, xpos, 4.50, 3.61, .76, detail, 18, MUTED)
    text(slide, .70, 5.83, 11.8, .34, "华南农业大学 · 数学与信息学院、软件学院 · 软件工程5班", 14, MUTED)
    text(slide, .70, 6.35, 11.8, .39, "第31组   组长：梁镇贤   汇报：刘亦达   ｜   2026年10月26日", 16, INK)

    slide = page(presentation, "基础 · 同源场景与一阶知识库", "同一份股权图，区分控制与权益", "依据：docs/lab02/理论与算法.md §2；core/kb/rules.py；outputs/lab02/demo/traces.json")
    text(slide, .70, 1.92, 12, .40, "Holds(x,c,r) 等谓词  ·  16实体基础 / 29实体扩展  ·  27条安全正Horn规则", 17, MUTED)
    nodes = [1.17, 3.37, 5.57, 7.77]
    for index, (xpos, label) in enumerate(zip(nodes, ["P1", "C1", "C2", "C3"])):
        box(slide, xpos, 2.83, .79, .79, ACCENT if label == "C3" else "E2D5C4", True)
        text(slide, xpos, 3.01, .79, .37, label, 21, BG if label == "C3" else INK, True, PP_ALIGN.CENTER)
        if index < 3:
            line(slide, xpos + .79, 3.22, nodes[index + 1], 3.22, ACCENT, True)
            text(slide, xpos + .86, 2.73, 1.27, .35, ["80%", "65%", "35%"][index], 18, ACCENT, True, PP_ALIGN.CENTER)
    line(slide, nodes[1] + .39, 3.62, nodes[1] + .39, 4.03, BLUE)
    line(slide, nodes[1] + .39, 4.03, nodes[3] + .39, 4.03, BLUE)
    line(slide, nodes[3] + .39, 4.03, nodes[3] + .39, 3.62, BLUE, True)
    text(slide, 4.73, 4.15, 2.78, .35, "C1直接持股20%", 16, BLUE, False, PP_ALIGN.CENTER)
    text(slide, 9.43, 2.67, 2.65, .95, "55%", 46, ACCENT, True)
    text(slide, 9.46, 3.55, 2.65, .77, "C1可支配C3\n35% + 20%", 18, MUTED)
    text(slide, .74, 4.92, 11.8, .41, "两段绝对控股 + 两份股权合计55% ⇒ C1控制C3 ⇒ P1控制C3", 19, INK, True)
    text(slide, .74, 5.64, 11.8, .49, "经济权益：0.8 × 0.2 + 0.8 × 0.65 × 0.35 = 34.2%", 23, BLUE)
    text(slide, .74, 6.36, 11.8, .34, "控制候选与简单路径权益回答不同问题；不是现实企业的法律认定。", 15, MUTED)

    slide = page(presentation, "基础 · 算法、证明与可靠性", "保留完整证明，也说明适用边界", "依据：experiments/lab02/deterministic/；docs/lab02/理论与算法.md §3–5")
    for xpos, title, formula, body in [(.72, "正向推理", "F(t+1) = T_R(F(t))", "事实驱动到不动点\n保存规则、绑定与前提\n索引与delta减少重复匹配"),
                                      (4.92, "表格化反向", "目标 → 子目标 → 答案表", "按查询组织推理\n复用中间答案、处理递归\n未证实与预算耗尽分开"),
                                      (9.08, "独立归结子案例", "KB ∧ ¬Q ⇒ □", "子句化、合一、归结\n证明步骤可独立重放\n不冒充完整算术规则库")]:
        text(slide, xpos, 2.17, 3.52, .49, title, 23, ACCENT, True)
        text(slide, xpos, 2.99, 3.53, .78, formula, 19, BLUE)
        text(slide, xpos, 3.90, 3.51, 1.54, body, 18, MUTED)
    line(slide, .72, 5.69, 12.60, 5.69)
    text(slide, .73, 6.01, 11.86, .72, "有限常量域 · 安全正Horn · 无函数项 · 无资源截断\n逻辑可靠 ≠ 资料真实；推不出 ≠ 已证伪。", 19, INK)

    slide = page(presentation, "基础 · 两种不确定性推理", "同一组指标，两类不确定性推理", "依据：uncertain/bayes.py、fuzzy.py；outputs/lab02/experiments/summary.csv（合成指标实验）")
    text(slide, .75, 1.93, 11.85, .41, "不透明度 / 关联交易比例 / 担保负荷 / 回路线索：同一输入、不同测度", 18, MUTED)
    line(slide, 6.58, 2.63, 6.58, 6.39)
    text(slide, .77, 2.64, 5.43, .50, "Bayes：条件概率", 26, BLUE, True)
    text(slide, .77, 3.38, 5.43, .81, "后验 ∝ 先验 × 观测指标似然\n先验、CPT由训练集估计", 21, INK)
    text(slide, .79, 4.48, 5.38, 1.17, "缺测：从观测证据中省略\n本项目星形模型评分为 O(d)\n不把条件独立假设当作事实", 18, MUTED)
    text(slide, .79, 5.93, 5.38, .44, f"清洁F1  {float(bayes['f1_mean']):.4f}   ·   {float(bayes['inference_us_mean']):.1f} μs/样本", 19, BLUE, True)
    text(slide, 7.02, 2.64, 5.30, .50, "Mamdani：模糊隶属度", 26, ACCENT, True)
    text(slide, 7.02, 3.38, 5.31, .81, "隶属度 → min/max聚合 → 质心\n训练集拟合15组参数候选", 21, INK)
    text(slide, 7.04, 4.48, 5.25, 1.17, "缺测：评分向中性值收缩\n规则数 × 输出网格带来开销\n模糊评分不冒称后验概率", 18, MUTED)
    text(slide, 7.04, 5.93, 5.25, .44, f"清洁F1  {float(fuzzy['f1_mean']):.4f}   ·   {float(fuzzy['inference_us_mean']):.1f} μs/样本", 19, ACCENT, True)
    text(slide, .78, 6.55, 11.8, .30, "参数与权重仅用训练/验证集选择；三种子均值不是现实企业准确率。", 14, MUTED)

    slide = page(presentation, "进阶 · 确定性优化", "闭包不变，减少重复匹配", "数据：outputs/lab02/benchmarks/forward_summary.csv；本机均值，非跨机器固定倍数")
    text(slide, .73, 1.93, 11.9, .42, "1 / 4 / 8 / 16组结构 × 三次重复；以下为16组：3031条事实，20轮迭代", 18, MUTED)
    chart = add_chart(slide, XL_CHART_TYPE.BAR_CLUSTERED, .73, 2.59, 7.77, 3.29,
                      ["全扫描", "索引朴素", "半朴素"], [("平均耗时/ms", [float(row['elapsed_ms_mean']) for row in benchmark])])
    chart.value_axis.minimum_scale, chart.value_axis.maximum_scale = 1, 100000
    scaling = chart.value_axis._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}scaling")
    log_base = OxmlElement("c:logBase")
    log_base.set("val", "10")
    scaling.insert(0, log_base)
    chart.plots[0].data_labels.number_format = "0.00"
    for point, color in zip(chart.series[0].points, [GRAY, BLUE, ACCENT]):
        point.format.fill.fore_color.rgb = RGBColor.from_string(color)
    text(slide, 9.06, 2.74, 3.4, .97, f"{speedup:.2f}×", 55, ACCENT, True)
    text(slide, 9.10, 3.88, 3.32, .86, "半朴素相对\n索引朴素加速", 21, INK)
    text(slide, 9.10, 5.03, 3.33, .83,
         f"候选检查\n{int(float(benchmark[1]['candidate_checks_mean'])):,} → {int(float(benchmark[2]['candidate_checks_mean'])):,}", 19, MUTED)
    ratios = [float(lookup(forward, groups=str(groups), model="indexed_naive")["elapsed_ms_mean"]) / float(lookup(forward, groups=str(groups), model="semi_naive")["elapsed_ms_mean"]) for groups in [1, 4, 8, 16]]
    text(slide, .80, 6.21, 11.7, .36, "不同规模加速比：" + "   /   ".join(f"{groups}组 {ratio:.2f}×" for groups, ratio in zip([1, 4, 8, 16], ratios)), 17, INK)
    text(slide, .80, 6.68, 11.7, .25, "横轴为对数耗时；先检查完整闭包一致，再解释时间、迭代与搜索空间。", 12, MUTED)

    slide = page(presentation, "进阶 · 参数学习与抗干扰", "更准不等于更稳：参数与扰动都要测", "数据：outputs/lab02/experiments/{summary,sensitivity,ablations}.csv；误差与全部条件见报告")
    text(slide, .74, 1.93, 11.9, .40, "3种子 · 8条件 · 每种子训练1000 / 验证300 / 测试600；模型在测试前冻结", 17, MUTED)
    conditions = ["clean", "noise_010", "noise_020", "noise_030", "severe"]
    chart = add_chart(slide, XL_CHART_TYPE.LINE_MARKERS, .75, 2.63, 7.17, 3.30,
                      ["清洁", "噪声.10", "噪声.20", "噪声.30", "重扰动"],
                      [(label, [float(lookup(scores, scenario=scenario, model=model)["f1_mean"]) for scenario in conditions])
                       for label, model in [("学习Bayes", "bayes_learned"), ("拟合Mamdani", "fuzzy_learned")]], False)
    chart.value_axis.minimum_scale, chart.value_axis.maximum_scale, chart.value_axis.major_unit = .5, 1, .1
    text(slide, 1.24, 2.38, 2.50, .27, "— 学习Bayes", 14, BLUE, True)
    text(slide, 4.25, 2.38, 3.05, .27, "— 拟合Mamdani", 14, ACCENT, True)
    text(slide, 8.53, 2.76, 3.9, .42, "CPT学习：清洁F1", 21, INK, True)
    text(slide, 8.53, 3.43, 3.9, .55, f"{float(bayes_manual['f1_mean']):.4f} → {float(bayes['f1_mean']):.4f}", 25, BLUE, True)
    text(slide, 8.55, 4.26, 3.83, 1.23, "拟合15组模糊参数\n同步检查参数敏感性\n零权重移除可无变化", 19, MUTED)
    severe_bayes = lookup(scores, scenario="severe", model="bayes_learned")
    text(slide, .84, 6.15, 11.5, .35,
         f"重扰动：噪声σ=0.35 + 缺测40%；学习Bayes F1降为{float(severe_bayes['f1_mean']):.4f}。", 18, INK)
    text(slide, .84, 6.65, 11.5, .29, "主图展示5个条件；缺测、混合与消融详表见报告。回退基线不等于组件边际贡献。", 12, MUTED)

    slide = page(presentation, "创新方向1 · 混合推理模型", "创新：把评分与最终行动分开", "依据：experiments/lab02/hybrid/{pipeline,model,decision,calibration}.py")
    for xpos, title, value in [(.77, "结构证明", "可重放的预警 A"), (5.02, "共同评分", "s = Σ wᵢsᵢ"), (9.22, "复核行动", "review_required")]:
        text(slide, xpos, 2.19, 3.22, .42, title, 24, ACCENT, True)
        text(slide, xpos, 2.95, 3.24, .42, value, 21, INK)
    line(slide, 3.96, 2.78, 4.72, 2.78, BLUE, True)
    line(slide, 8.19, 2.78, 8.94, 2.78, BLUE, True)
    table(slide, .78, 3.90, [2.33, 5.50, 3.90], [
        ["策略", "最终行动", "原始评分与证明"],
        ["不锁定", "d_U = 1[s ≥ 0.5]", "始终保留"],
        ["结构硬锁定", "d_L = d_U ∨ A", "与不锁定相同"],
        ["校准复核", "d_C = d_U ∨ (A ∧ s ≥ η)", "仅改变最终行动"],
    ], .54, 18)
    text(slide, .83, 6.43, 11.6, .35, "A：观测图推出的高预警；η：仅由验证扰动选择的最低评分支持。", 16, MUTED)

    slide = page(presentation, "创新验证 · 图输入到最终复核", "误报减少，但不能宣称全面胜出", "数据：outputs/lab02/end_to_end/summary.csv；三个种子均值，无统计显著性宣称")
    text(slide, .74, 1.93, 11.9, .41, "3种子 × 6条件 × 11方法；以下为20%错误闭环、同一自适应评分模型", 18, MUTED)
    chart = add_chart(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, .79, 2.64, 7.56, 3.41,
                      ["不锁定", "硬锁定", "校准复核"], [("误报率", [float(row["false_positive_rate_mean"]) for row in graph_rows])])
    chart.value_axis.minimum_scale, chart.value_axis.maximum_scale, chart.value_axis.major_unit = 0, .6, .2
    chart.value_axis.tick_labels.number_format = "0%"
    chart.plots[0].data_labels.number_format = "0.00%"
    for point, color in zip(chart.series[0].points, [BLUE, GRAY, ACCENT]):
        point.format.fill.fore_color.rgb = RGBColor.from_string(color)
    text(slide, 8.94, 2.81, 3.63, 1.02, f"−{false_positive_drop:.2f}", 51, ACCENT, True)
    text(slide, 9.02, 4.00, 3.54, .82, "个百分点\n校准相对硬锁定误报率", 19, INK)
    text(slide, 9.02, 5.08, 3.55, 1.06, "分数相同，行动不同\n硬锁定为0的约束违反率\n不是完美风险预测", 16, MUTED)
    text(slide, .85, 6.25, 11.5, .36, "F1：" + "  /  ".join(f"{label} {float(row['f1_mean']):.4f}" for label, row in zip(["不锁定", "硬锁定", "校准"], graph_rows)), 18, INK, True)
    text(slide, .85, 6.70, 11.5, .25, "降低硬锁定代价，不等于全面超过不锁定；合成高复核需求标签并非真实企业风险。", 12, MUTED)

    slide = page(presentation, "结论 · 五维比较与复现边界", "方法各有所长，结论必须可复现", "依据：docs/lab02/理论与算法.md §10；outputs/qa/lab02-fixes/completion.json；完整方法对照见附录")
    table(slide, .77, 1.98, [1.52, 5.01, 5.27], [
        ["维度", "确定性推理", "不确定性推理"],
        ["建模难度", "需明确谓词、规则与守卫", "需数据分布或隶属度参数"],
        ["效率", "闭包与连接规模决定开销", "星形Bayes轻；模糊网格更重"],
        ["容错性", "缺事实可能中断证明", "可处理不完整指标，仍会退化"],
        ["可信度", "路径可追溯，不保证资料真实", "概率/模糊语义不同，需校准"],
        ["落地限制", "需核验规则与结构事实", "需验证数据分布与决策代价"],
    ], .52, 16)
    text(slide, .82, 5.48, 11.55, .55, f"{TEST_COUNT}项项目测试  ·  23,760条图级决策回放", 27, ACCENT, True)
    text(slide, .84, 6.23, 11.54, .55, "数据与标签均为教学构造。下一步：资料核验、跨分布验证与人工复核。", 18, INK)

    slide = page(presentation, "附录 A · 方法比较", "五个维度，不把不同测度混为一谈", "依据：docs/lab02/理论与算法.md §10；正式报告“对比分析与讨论”")
    table(slide, .69, 1.96, [1.28, 3.61, 3.53, 3.50], [
        ["维度", "确定性规则/归结", "星形Bayes", "Mamdani"],
        ["建模", "谓词、Horn规则、数值守卫", "先验、分箱、CPT、独立假设", "隶属函数、规则、输出集合"],
        ["效率", "连接/闭包/搜索可能膨胀", "d项观测的轻量评分", "规则激活与网格聚合"],
        ["容错", "环可处理；缺事实难补证", "缺项可省略；噪声仍误导", "平滑边界；缺测向中性收缩"],
        ["可信", "完整证明可重放", "可用Brier检查概率质量", "隶属度不是后验概率"],
        ["落地", "教学规则需现实核验", "训练偏移/依赖性需复查", "经验参数需校准和迁移验证"],
    ], .68, 15)
    text(slide, .76, 6.39, 11.8, .42, "时间、精度、容错、可信度分开比较；没有一种方法在所有条件下全面占优。", 19, INK, True)

    slide = page(presentation, "附录 B · 敏感性与消融", "展示参数影响，也保留没有增益的结果", "数据：outputs/lab02/experiments/sensitivity.csv、ablations.csv；三种子均值")
    grouped = defaultdict(list)
    for row in sensitivity:
        grouped[(row["model"], row["parameter"], float(row["value"]))].append(float(row["f1"]))
    for xpos, model, parameter, title in [(.77, "bayes", "prior", "Bayes先验敏感性"), (6.96, "fuzzy", "high_start", "模糊高隶属起点敏感性")]:
        points = sorted((key[2], sum(values) / len(values)) for key, values in grouped.items() if key[:2] == (model, parameter))
        text(slide, xpos, 1.97, 5.57, .39, title, 20, BLUE if model == "bayes" else ACCENT, True)
        chart = add_chart(slide, XL_CHART_TYPE.LINE_MARKERS, xpos, 2.59, 5.46, 2.58,
                          [f"{value:g}" for value, _ in points], [("F1", [value for _, value in points])], False)
        chart.value_axis.minimum_scale, chart.value_axis.maximum_scale, chart.value_axis.major_unit = .6, 1, .1
    unchanged = sum(row["kind"] == "unchanged_zero_weight" for row in ablations)
    fallback = sum(row["kind"] == "baseline_fallback" for row in ablations)
    text(slide, .83, 5.53, 11.55, .47, f"{len(ablations)}条消融操作：{unchanged}条移除零权重后不变；{fallback}条为显式回退基线。", 22, ACCENT, True)
    text(slide, .85, 6.18, 11.52, .65, "清洁指标实验中调优权重退化为纯Bayes，不能宣称融合增益。\n敏感性扫描仅作诊断，不按测试F1重新选参；D-S冲突折扣为额外扩展。", 17, MUTED)

    slide = page(presentation, "附录 C · 理论与完整证据", "从公式、规则到可重放的证明", "依据：docs/lab02/理论与算法.md；outputs/lab02/demo/traces.json；deterministic/resolution.py")
    formulas = [
        ("一阶控制与不动点", "Holds(x,c,r) ∧ r > 0.5 ∧ x ≠ c ⇒ DirectControls(x,c)", "F(t+1) = T_R(F(t))；控制传递与份额合计使用明确规则，不等于经济权益。"),
        ("Bayes与Mamdani", "P(R=1 | b_O) ∝ π · ∏(j∈O) P(b_j | R=1)", "模糊质心：s = Σ(z · μ聚合(z)) / Σμ聚合(z)；未观测项不冒充零证据。"),
        ("归结与证明范围", "KB ∧ ¬Q → 子句化 → 合一/归结 → 空子句 □", "独立股权子案例；完整事实、规则ID、绑定与步骤见保存的trace，不扩大可靠性声明。"),
    ]
    for index, (title, formula, explanation) in enumerate(formulas):
        ypos = 1.98 + index * 1.57
        text(slide, .77, ypos, 11.82, .36, title, 19, ACCENT, True)
        text(slide, .80, ypos + .48, 11.80, .48, formula, 22, INK)
        text(slide, .81, ypos + 1.05, 11.75, .39, explanation, 15, MUTED)

    slide = page(presentation, "附录 D · 系统复现与成员分工", "系统复现与成员分工", "依据：项目README、成员责任分工表。")
    rows = [["姓名", "学号", "职能", "责任范围"]]
    for member in brief["members"]:
        rows.append([member["name"], member["id"], member["role"], member["responsibility"]])
    table(slide, .69, 1.97, [1.10, 1.92, 2.10, 6.80], rows, .58, 13)
    text(slide, .76, 5.72, 11.81, .34, "系统入口：打开网站.cmd / 启动实验二.cmd；环境与运行说明见报告附录 B。", 16, INK)
    text(slide, .76, 6.20, 11.82, .39, "核心源码：core/kb/ 与 experiments/lab02/", 19, BLUE)
    text(slide, .76, 6.68, 11.81, .24, "实验成果：科研型报告、推理源码、实验数据、结果图表与创新说明。", 12, MUTED)

    for index, slide in enumerate(presentation.slides):
        if index < len(narration["slides"]):
            entry = narration["slides"][index]
            body = "".join(cue.get("speech", cue["text"]) for cue in entry["cues"])
            slide.notes_slide.notes_text_frame.text = f"实验二汇报 · 第{index+1}页 · 汇报：刘亦达\n{body}"
        else:
            slide.notes_slide.notes_text_frame.text = "补充实验材料：方法比较、敏感性与消融、理论公式及成员分工。"
        transition = OxmlElement("p:transition")
        transition.set("advClick", "1")
        slide._element.append(transition)
    main_count = len(narration["slides"])
    original_deck = BytesIO()
    presentation.save(original_deck)

    def save_subset(indices, destination):
        subset = Presentation(BytesIO(original_deck.getvalue()))
        for index in reversed(range(len(subset.slides))):
            if index not in indices:
                slide_id = subset.slides._sldIdLst[index]
                subset.part.drop_rel(slide_id.rId)
                subset.slides._sldIdLst.remove(slide_id)
        if destination == SUPPLEMENTARY:
            subset.core_properties.title += "：补充材料"
        subset.save(destination)
        return sum(shape.has_chart for slide in subset.slides for shape in slide.shapes)

    main_charts = save_subset(range(main_count), PPTX)
    supplementary_charts = save_subset(range(main_count, len(presentation.slides)), SUPPLEMENTARY)
    assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() == original_hash
    sources = [forward_path, score_path, graph_path, sensitivity_path, ablation_path]
    manifest = {"pptx": str(PPTX.relative_to(ROOT)), "main_slides": main_count, "backup_slides": 0,
                "supplementary_pptx": str(SUPPLEMENTARY.relative_to(ROOT)),
                "supplementary_slides": len(presentation.slides) - main_count,
                "native_editable_charts": main_charts,
                "supplementary_editable_charts": supplementary_charts,
                "visible_source_footers": False,
                "original_pptx_sha256": original_hash, "original_pptx_unchanged": True,
                "data_root": str(DATA_ROOT), "project_tests": TEST_COUNT,
                "sources": {name: hashlib.sha256((DATA_ROOT / name).read_bytes()).hexdigest() for name in sources},
                "speedup": speedup, "false_positive_drop_pp": false_positive_drop}
    (WORK / "deck-build.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2), flush=True)


def export_deck(pptx, references, report_name):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    app = deck = None
    pptx_hash = hashlib.sha256(pptx.read_bytes()).hexdigest()
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        deck = app.Presentations.Open(str(pptx), True, False, False)
        deck.SaveAs(str(pptx.with_suffix(".pdf")), 32)
        references.mkdir(exist_ok=True)
        flags = []
        for slide_index in range(1, deck.Slides.Count + 1):
            slide = deck.Slides(slide_index)
            slide.Export(str(references / f"slide-{slide_index:02}.png"), "PNG", 1600, 900)
            for shape_index in range(1, slide.Shapes.Count + 1):
                shape = slide.Shapes(shape_index)
                try:
                    if shape.HasTextFrame and shape.TextFrame.HasText:
                        content = shape.TextFrame.TextRange
                        if content.BoundHeight > shape.Height + 4:
                            flags.append({"slide": slide_index, "shape": shape_index, "text": content.Text,
                                          "height": shape.Height, "text_height": content.BoundHeight})
                except Exception:
                    pass
        result = {"slide_count": int(deck.Slides.Count), "overflow_flags": flags,
                  "pdf": str(pptx.with_suffix(".pdf").relative_to(ROOT)), "pptx_unchanged_by_export": hashlib.sha256(pptx.read_bytes()).hexdigest() == pptx_hash}
        (WORK / report_name).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
    finally:
        if deck is not None:
            deck.Close()
        if app is not None:
            try:
                if app.Presentations.Count == 0:
                    app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


def export():
    export_deck(PPTX, WORK / "slide-renders", "deck-export.json")
    export_deck(SUPPLEMENTARY, WORK / "supplementary-renders", "supplementary-export.json")


def main():
    global WORK, DELIVERY, PPTX, SUPPLEMENTARY, DATA_ROOT, TEST_COUNT
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stage", choices=["build", "export", "all"], default="all")
    parser.add_argument("--source-root", type=Path, default=DATA_ROOT)
    parser.add_argument("--work-dir", type=Path, default=WORK)
    parser.add_argument("--output-dir", type=Path, default=DELIVERY)
    parser.add_argument("--stem", default=PPTX.stem)
    parser.add_argument("--test-count", type=int, default=TEST_COUNT)
    args = parser.parse_args()
    if Path(args.stem).name != args.stem or any(c in args.stem for c in '/\\:'):
        parser.error("The deck stem must be a filename, not a path")
    if args.test_count <= 0:
        parser.error("The project test count must be positive")
    DATA_ROOT, WORK, DELIVERY = (p.resolve() for p in [args.source_root, args.work_dir, args.output_dir])
    WORK.relative_to(ROOT)
    DELIVERY.relative_to(ROOT)
    PPTX = DELIVERY / (args.stem + ".pptx")
    SUPPLEMENTARY = DELIVERY / "第31组_推理实验答辩_补充材料.pptx"
    TEST_COUNT = args.test_count
    if args.stage in ("build", "all"):
        build()
    if args.stage in ("export", "all"):
        export()


if __name__ == "__main__":
    main()
