"""Record the original lab01 and lab02 PowerPoint decks as one narrated video.

The two source decks are opened read-only. They are recorded separately in
PowerPoint and joined at the narration boundary, so no application switch is
included in the video. Source slides, synchronized subtitles, speech, and an
editable narration document are kept with a reproducible timeline.
"""
from __future__ import annotations

import argparse
import asyncio
import copy
import csv
import hashlib
import json
import os
from pathlib import Path
import shutil
import sys
import time

import create_lab02_ppt_video as media

ORIGINAL_WRITE_SUBTITLES = media.write_subtitles
ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "outputs/qa/lab01-lab02-joint-20261027"
DELIVERY = ROOT / "deliverables/lab01-lab02"
SUBMISSION = ROOT.parent / "人工智能实验2"
STEM = "第31组_实验一与实验二_联合汇报_4分钟内"
VIDEO = DELIVERY / f"{STEM}.mp4"
REFERENCES = WORK / "slide-renders"
NARRATION_SOURCE = ROOT / "deliverables/lab01/最终演讲稿.docx"
SOURCES = [
    {"lab": "lab01", "name": "实验一", "count": 15,
     "pptx": ROOT / "deliverables/lab01/企业股权知识表示与风险推理.pptx"},
    {"lab": "lab02", "name": "实验二", "count": 9,
     "pptx": SUBMISSION / "第31组_推理实验汇报.pptx"},
]


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_subtitles(timeline):
    ORIGINAL_WRITE_SUBTITLES(timeline)
    subtitle_file = WORK / "subtitles.ass"
    text = subtitle_file.read_text(encoding="utf-8-sig")
    # Lab01 has a header starting at y=46. Use the empty strip above it.
    text = text.replace("Narration,Microsoft YaHei,42,", "Narration,Microsoft YaHei,38,")
    text = text.replace(",110,110,24,1", ",110,110,8,1")
    subtitle_file.write_text(text, encoding="utf-8-sig")


def configure(work=WORK, pptx=None):
    work.mkdir(parents=True, exist_ok=True)
    DELIVERY.mkdir(parents=True, exist_ok=True)
    media.WORK = work
    media.DELIVERY = DELIVERY
    media.PPTX = pptx or SOURCES[0]["pptx"]
    media.REFERENCES = REFERENCES
    media.STEM = STEM
    media.VIDEO = VIDEO
    media.SRT = DELIVERY / f"{STEM}.srt"
    media.write_subtitles = write_subtitles
    # The reusable helper's guide is specific to the former nine-page video.
    media.write_editing_guide = lambda *_: None


def protect_sources():
    baseline_path = WORK / "source-baseline.json"
    if not baseline_path.exists():
        protected = [s["pptx"] for s in SOURCES] + [NARRATION_SOURCE]
        protected += [p for p in SUBMISSION.iterdir() if p.is_file() and p.name != VIDEO.name]
        baseline = {str(p): {"sha256": digest(p), "bytes": p.stat().st_size} for p in protected}
        write_json(baseline_path, baseline)
    baseline = read_json(baseline_path)
    for filename, expected in baseline.items():
        path = Path(filename)
        if not path.exists() or digest(path) != expected["sha256"]:
            raise RuntimeError(f"Protected source changed: {path}")
    return baseline


def export_sources():
    import pythoncom
    import win32com.client
    from PIL import Image, ImageDraw, ImageFont

    protect_sources()
    REFERENCES.mkdir(parents=True, exist_ok=True)
    pythoncom.CoInitialize()
    app = deck = None
    open_decks = []
    exported = []
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        for source in SOURCES:
            deck = app.Presentations.Open(str(source["pptx"]), True, False, False)
            open_decks.append(deck)
            if deck.Slides.Count != source["count"]:
                raise RuntimeError(f"Unexpected slide count: {source['pptx']}")
            aspect = float(deck.PageSetup.SlideWidth) / float(deck.PageSetup.SlideHeight)
            if abs(aspect - 16 / 9) > .001:
                raise RuntimeError("Source deck is not widescreen")
            for number in range(1, source["count"] + 1):
                index = len(exported) + 1
                target = REFERENCES / f"slide-{index:02}.png"
                slide = deck.Slides(number)
                slide.Export(str(target), "PNG", 1920, 1080)
                exported.append({"slide": index, "lab": source["lab"], "source_slide": number,
                                 "source_pptx": str(source["pptx"]), "image": str(target)})
                media.log(f"Exported {source['lab']} slide {number}/{source['count']}")
    finally:
        for opened in reversed(open_decks):
            try:
                opened.Close()
            except Exception:
                pass
        if app is not None:
            try:
                if app.Presentations.Count == 0:
                    app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()
    protect_sources()
    write_json(WORK / "source-slides.json", exported)
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 20)
    for source in SOURCES:
        entries = [item for item in exported if item["lab"] == source["lab"]]
        sheet = Image.new("RGB", (1920, ((len(entries) + 2) // 3) * 394), "#e4e1db")
        draw = ImageDraw.Draw(sheet)
        for index, entry in enumerate(entries):
            x, y = index % 3 * 640, index // 3 * 394
            draw.text((x + 12, y + 3), f"{source['name']}  {entry['source_slide']:02}", font=font, fill="#222222")
            with Image.open(entry["image"]) as picture:
                sheet.paste(picture.resize((640, 360)), (x, y + 34))
        sheet.save(WORK / f"source-contact-{source['lab']}.png")


def write_guide(timeline, settings):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    title = "实验一与实验二联合汇报逐页讲稿"
    intro = (f"成片 {timeline['duration']:.0f} 秒，1920×1080，30 帧/秒。"
             "按实验一15页、实验二9页的原顺序讲解，包含选题介绍和实验间过渡。"
             "使用普通话合成旁白，中文字幕已烧录到画面；独立SRT供剪辑使用。")
    lines = [f"# {title}", "", intro, "", "## 使用的原材料", ""]
    lines.extend(f"- {s['name']}：{s['pptx']}" for s in SOURCES)
    lines.extend([f"- 实验一原讲稿：{NARRATION_SOURCE}", "", "## 逐页时间表", "",
                  "| 视频页 | 原PPT | 入点 | 出点 | 主题 |", "|---|---|---|---|---|"])
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    document.add_heading(title, 0)
    document.add_paragraph(intro)
    document.add_paragraph("来源：实验一使用指定PPT和《最终演讲稿》；实验二使用桌面最终提交版《第31组_推理实验汇报》。")
    rows = []
    for slide, source in zip(timeline["slides"], settings["slides"]):
        label = "实验一" if source["lab"] == "lab01" else "实验二"
        start, end = media.timestamp(slide["start"]), media.timestamp(slide["end"])
        lines.append(f"| {slide['slide']} | {label}第{source['source_slide']}页 | {start} | {end} | {slide['title']} |")
        rows.append([slide["slide"], source["lab"], source["source_slide"], slide["start"], slide["end"], start, end, slide["title"]])
    for slide, source in zip(timeline["slides"], settings["slides"]):
        label = "实验一" if source["lab"] == "lab01" else "实验二"
        heading = f"{label} · 第{source['source_slide']}页 · {slide['title']}"
        times = f"{media.timestamp(slide['start'])} — {media.timestamp(slide['end'])}"
        text = "".join(cue["text"] for cue in source["cues"])
        lines.extend(["", f"## {heading}", "", times, "", text])
        paragraph = document.add_heading(heading, 2)
        paragraph.paragraph_format.keep_with_next = True
        document.add_paragraph(times)
        document.add_paragraph(text)
    lines.extend(["", "## 说明", "", "两份原PPT和实验一原讲稿未修改。视频中的测试与实验数值引用原PPT已有记录，本次仅重新制作联合汇报视频。", ""])
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("联合汇报讲稿 · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(DELIVERY / f"{STEM}_逐页讲稿.docx")
    (DELIVERY / f"{STEM}_逐页讲稿.md").write_text("\n".join(lines), encoding="utf-8-sig")
    with (DELIVERY / f"{STEM}_剪辑时间表.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["video_slide", "lab", "source_slide", "start_seconds", "end_seconds", "start_timecode", "end_timecode", "title"])
        writer.writerows(rows)


def prepare():
    protect_sources()
    settings = read_json(WORK / "narration.json")
    exported = read_json(WORK / "source-slides.json")
    expected = [(entry["lab"], entry["source_slide"]) for entry in exported]
    actual = [(entry["lab"], entry["source_slide"]) for entry in settings["slides"]]
    if actual != expected or len(actual) != 24:
        raise RuntimeError("Narration must cover all original slides in their original order")
    asyncio.run(media.prepare())
    timeline = read_json(WORK / "timeline.json")
    timeline["source_decks"] = [{"lab": s["lab"], "pptx": str(s["pptx"]), "sha256": digest(s["pptx"]),
                                 "slide_count": s["count"]} for s in SOURCES]
    for entry, source in zip(timeline["slides"], settings["slides"]):
        entry.update({"lab": source["lab"], "source_slide": source["source_slide"]})
    write_json(WORK / "timeline.json", timeline)
    for source in SOURCES:
        subset = copy.deepcopy([s for s in timeline["slides"] if s["lab"] == source["lab"]])
        start = subset[0]["start"]
        duration = subset[-1]["end"] - start
        local = {**timeline, "slides": subset, "duration": duration, "target_duration": duration,
                 "source_pptx": str(source["pptx"]), "source_sha256": digest(source["pptx"]),
                 "global_start": start}
        for slide in subset:
            slide["video_slide"] = slide["slide"]
            slide["slide"] = slide["source_slide"]
            for field in ["start", "end", "speech_start"]:
                slide[field] -= start
            for cue in slide["cues"]:
                cue["start"] -= start
                cue["end"] -= start
        folder = WORK / f"capture-{source['lab']}"
        folder.mkdir(exist_ok=True)
        write_json(folder / "timeline.json", local)
    write_guide(timeline, settings)
    print(json.dumps({"duration": timeline["duration"], "tempo_factor": timeline["tempo_factor"],
                      "lab01_end": timeline["slides"][14]["end"], "slides": len(timeline["slides"])}, ensure_ascii=False), flush=True)


def record(lab=None, probe=False):
    protect_sources()
    for source in SOURCES:
        if lab and source["lab"] != lab:
            continue
        configure(WORK / f"capture-{source['lab']}", source["pptx"])
        media.record(probe=probe)
    configure()
    protect_sources()


def measure_captures():
    import cv2
    import numpy as np

    report = []
    for source in SOURCES:
        folder = WORK / f"capture-{source['lab']}"
        local = read_json(folder / "timeline.json")
        capture = read_json(folder / "capture-events.json")
        cap = cv2.VideoCapture(str(folder / "slideshow-raw.mkv"))
        references = []
        for slide in local["slides"]:
            reference = cv2.imread(str(REFERENCES / f"slide-{slide['video_slide']:02}.png"))
            if reference is None:
                raise RuntimeError("Missing source reference")
            references.append(cv2.resize(reference, (160, 90)).astype(np.float32))

        def classify(frame):
            small = cv2.resize(frame, (160, 90)).astype(np.float32)
            errors = [float(np.mean((small - ref) ** 2)) for ref in references]
            best = int(np.argmin(errors))
            return best, errors[best]

        measured = []
        for index, slide in enumerate(local["slides"][1:], 1):
            begin = max(0, slide["start"] + capture["warmup"] - 1.3)
            cap.set(cv2.CAP_PROP_POS_MSEC, begin * 1000)
            transition = None
            seen_previous = False
            for _ in range(90):
                ok, frame = cap.read()
                if not ok:
                    break
                second = cap.get(cv2.CAP_PROP_POS_MSEC) / 1000
                best, error = classify(frame)
                if best == index - 1 and error < 1200:
                    seen_previous = True
                if best == index and error < 500 and seen_previous:
                    transition = second
                    break
            if transition is None:
                raise RuntimeError(f"Cannot locate captured transition to {source['lab']} slide {slide['slide']}")
            measured.append({"slide": slide["slide"], "raw_transition": transition,
                             "planned_transition": slide["start"], "offset": transition - slide["start"]})
        offsets = [entry["offset"] for entry in measured]
        offset = float(np.median(offsets))
        trim_frames = round(offset * media.FPS)
        jitter = max(abs(entry["offset"] - trim_frames / media.FPS) for entry in measured)
        if jitter > .25 or not 0 <= offset <= 2:
            raise RuntimeError(f"Capture timing out of bounds: {source['lab']}, offset={offset}, jitter={jitter}")
        capture.update({"measured_trim": trim_frames / media.FPS, "trim_frames": trim_frames,
                        "maximum_transition_error": jitter, "measured_transitions": measured})
        write_json(folder / "capture-events.json", capture)
        report.append({"lab": source["lab"], "trim_frames": trim_frames,
                       "maximum_transition_error": jitter, "duration": local["duration"]})
        cap.release()
    write_json(WORK / "capture-alignment.json", report)
    media.log(json.dumps(report, ensure_ascii=False, indent=2))


def render():
    measure_captures()
    timeline = read_json(WORK / "timeline.json")
    arguments = []
    filters = []
    for index, source in enumerate(SOURCES):
        folder = WORK / f"capture-{source['lab']}"
        local = read_json(folder / "timeline.json")
        capture = read_json(folder / "capture-events.json")
        start_frame = capture["trim_frames"]
        frame_count = sum(slide["frame_count"] for slide in local["slides"])
        arguments += ["-i", folder / "slideshow-raw.mkv"]
        filters.append(f"[{index}:v]trim=start_frame={start_frame}:end_frame={start_frame + frame_count},setpts=PTS-STARTPTS[v{index}]")
    arguments += ["-i", WORK / "narration.wav"]
    subtitle_file = (WORK / "subtitles.ass").relative_to(ROOT).as_posix()
    filters.append(f"[v0][v1]concat=n=2:v=1:a=0,ass=filename='{subtitle_file}'[video]")
    temporary = WORK / "final-pending.mp4"
    arguments += ["-filter_complex", ";".join(filters), "-map", "[video]", "-map", "2:a:0",
                  "-t", timeline["duration"], "-r", media.FPS,
                  "-c:v", "libx264", "-preset", "medium", "-crf", "18", "-pix_fmt", "yuv420p",
                  "-af", "loudnorm=I=-16:TP=-1.5:LRA=7", "-ar", "48000", "-ac", "2",
                  "-c:a", "aac", "-b:a", "192k", "-movflags", "+faststart",
                  "-metadata", f"title={STEM}", "-metadata", "comment=两套原PPT逐页放映录制；普通话合成旁白；中文字幕内嵌",
                  temporary]
    media.run_ffmpeg(arguments, "render.log")
    if VIDEO.exists():
        VIDEO.replace(WORK / f"previous-{time.strftime('%Y%m%d-%H%M%S')}.mp4")
    temporary.replace(VIDEO)
    media.log(f"Rendered {VIDEO}")


def verify():
    protect_sources()
    media.verify()
    report_path = WORK / "verification.json"
    report = read_json(report_path)
    timeline = read_json(WORK / "timeline.json")
    settings = read_json(WORK / "narration.json")
    for slide, script in zip(timeline["slides"], settings["slides"]):
        if [cue["text"] for cue in slide["cues"]] != [cue["text"] for cue in script["cues"]]:
            raise RuntimeError("Final subtitles differ from the approved narration data")
        for cue in slide["cues"]:
            if not slide["start"] <= cue["start"] < cue["end"] <= slide["end"]:
                raise RuntimeError("A subtitle is outside its source slide")
    report.update({"source_decks": timeline["source_decks"], "original_sources_unchanged": True,
                   "lab01_slides": 15, "lab02_slides": 9, "lab01_end_seconds": timeline["slides"][14]["end"],
                   "video_sha256": digest(VIDEO), "capture_alignment": read_json(WORK / "capture-alignment.json")})
    write_json(report_path, report)
    protect_sources()


def deliver():
    protect_sources()
    report = read_json(WORK / "verification.json")
    if digest(VIDEO) != report["video_sha256"]:
        raise RuntimeError("Only the verified video may be copied to the submission directory")
    destination = SUBMISSION / VIDEO.name
    if destination.exists() and digest(destination) != digest(VIDEO):
        shutil.copy2(destination, WORK / f"previous-submission-{time.strftime('%Y%m%d-%H%M%S')}.mp4")
    shutil.copy2(VIDEO, destination)
    if digest(destination) != digest(VIDEO):
        raise RuntimeError("Video copy did not verify")
    protect_sources()
    write_json(WORK / "delivery.json", {"video": str(destination), "sha256": digest(destination),
                                       "duration_seconds": report["duration_seconds"],
                                       "source_materials_unchanged": True})
    media.log(f"Delivered {destination}")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    os.chdir(ROOT)
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stage", required=True, choices=["export", "prepare", "probe", "record", "render", "verify", "deliver"])
    parser.add_argument("--lab", choices=["lab01", "lab02"])
    args = parser.parse_args()
    configure()
    if args.stage == "export":
        export_sources()
    elif args.stage == "prepare":
        prepare()
    elif args.stage in ["probe", "record"]:
        record(args.lab, probe=args.stage == "probe")
    else:
        {"render": render, "verify": verify, "deliver": deliver}[args.stage]()


if __name__ == "__main__":
    main()
