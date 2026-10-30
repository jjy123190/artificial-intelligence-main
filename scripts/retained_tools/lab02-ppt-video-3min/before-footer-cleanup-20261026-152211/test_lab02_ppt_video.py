"""Offline regression checks for the lab02 presentation/video tooling."""
import hashlib
import json
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts import create_lab02_ppt_video as video
from pptx import Presentation

WORK = ROOT / "outputs/qa/lab02-ppt-video-3min"


class TimingTests(unittest.TestCase):
    def test_exact_frame_allocation(self):
        frames = video.allocate_slide_frames([90, 120, 60], 300)
        self.assertEqual(sum(frames), 300)
        self.assertTrue(all(value >= minimum for value, minimum in zip(frames, [90, 120, 60])))

    def test_insufficient_budget_rejected(self):
        with self.assertRaises(ValueError):
            video.allocate_slide_frames([100, 200], 299)

    def test_invalid_or_empty_allocation_rejected(self):
        for counts in [[], [0], [-2, 5]]:
            with self.assertRaises(ValueError):
                video.allocate_slide_frames(counts, 10)

    def test_legacy_natural_allocation(self):
        self.assertEqual(video.allocate_slide_frames([90, 120]), [90, 120])

    def test_timecode_carry(self):
        self.assertEqual(video.timestamp(59.9996), "00:01:00,000")
        self.assertEqual(video.timestamp(180), "00:03:00,000")
        self.assertEqual(video.timestamp(179.9951, ass=True), "0:03:00.00")

    def test_speech_normalization(self):
        self.assertEqual(video.normalized("Ｆ1，C３：55％"), "f1c355")

    def test_legacy_defaults_stay_separate(self):
        self.assertEqual(video.WORK.name, "lab02-ppt-video")
        self.assertNotIn("3分钟", video.PPTX.name)


class PreparedAssetTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.timeline = json.loads((WORK / "timeline.json").read_text(encoding="utf-8"))
        cls.build = json.loads((WORK / "deck-build.json").read_text(encoding="utf-8"))
        cls.deck = Presentation(ROOT / cls.build["pptx"])

    def test_exact_three_minutes(self):
        self.assertEqual(self.timeline["duration"], 180)
        self.assertEqual(sum(slide["frame_count"] for slide in self.timeline["slides"]), 5400)
        self.assertEqual(len(self.timeline["slides"]), 9)

    def test_contiguous_slides_and_cues(self):
        previous_end = 0
        previous_caption_end = -1
        for slide in self.timeline["slides"]:
            self.assertAlmostEqual(slide["start"], previous_end)
            self.assertLess(slide["speech_start"] + slide["speech_duration"], slide["end"])
            for cue in slide["cues"]:
                self.assertGreaterEqual(cue["start"], slide["start"])
                self.assertLess(cue["start"], cue["end"])
                self.assertLessEqual(cue["end"], slide["end"])
                self.assertGreater(cue["start"], previous_caption_end)
                self.assertLessEqual(len(cue["text"]), 36)
                previous_caption_end = cue["end"]
            previous_end = slide["end"]

    def test_source_pptx_matches_timeline(self):
        actual = hashlib.sha256((ROOT / self.timeline["source_pptx"]).read_bytes()).hexdigest()
        self.assertEqual(actual, self.timeline["source_sha256"])

    def test_original_pptx_protected(self):
        original = ROOT / "deliverables/lab02/第31组_推理实验答辩.pptx"
        self.assertEqual(hashlib.sha256(original.read_bytes()).hexdigest(), self.build["original_pptx_sha256"])

    def test_editable_slides_notes_and_charts(self):
        self.assertEqual(len(self.deck.slides), 13)
        self.assertEqual(sum(shape.has_chart for slide in self.deck.slides for shape in slide.shapes), 5)
        for slide in self.deck.slides:
            self.assertTrue(slide.notes_slide.notes_text_frame.text)
            for element in slide._element:
                self.assertIsNone(element.get("advTm"))

    def test_presentation_excludes_production_instructions(self):
        forbidden = ["三分钟", "3分钟", "180秒", "97秒", "录屏", "新增录屏工具", "本版覆盖核查",
                     "经用户确认", "平台上传", "详细时间见同目录", "不计入", "制作信息"]
        for index, slide in enumerate(self.deck.slides, 1):
            chunks = [slide.notes_slide.notes_text_frame.text]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    chunks.append(shape.text)
                if shape.has_table:
                    chunks.extend(cell.text for row in shape.table.rows for cell in row.cells)
            combined = "\n".join(chunks)
            for phrase in forbidden:
                self.assertNotIn(phrase, combined, f"Unrelated production instruction on slide {index}")
        self.assertNotIn("分钟", self.deck.core_properties.title)

    def test_visible_caption_style_has_no_panel(self):
        subtitles = (WORK / "subtitles.ass").read_text(encoding="utf-8-sig")
        self.assertIn("100,100,0.3,0,1,2,0.5,8,110,110,24,1", subtitles)
        self.assertEqual(subtitles.count("Dialogue:"), sum(len(slide["cues"]) for slide in self.timeline["slides"]))

    def test_dataset_hashes_still_match(self):
        for relative, expected in self.build["sources"].items():
            self.assertEqual(hashlib.sha256((ROOT / relative).read_bytes()).hexdigest(), expected)


if __name__ == "__main__":
    unittest.main(verbosity=2)
