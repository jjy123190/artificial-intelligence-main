"""Validate the final report, then copy only DOCX and PDF into the submission parent."""
import csv
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
import time
import zipfile

from docx import Document
import fitz

WORK = Path(__file__).resolve().parent
PROJECT = WORK.parents[2]
DESTINATION = Path(r"C:\Users\Administrator\Desktop\人工智能实验2")
RELEASE = DESTINATION / "第31组_实验二源码"
STEM = "第31组_人工智能推理实验报告"
snapshot = json.loads((WORK / "evidence.json").read_text(encoding="utf-8"))
office = json.loads((WORK / "office-export.json").read_text(encoding="utf-8"))
assert 5000 <= office["main_text_word_count"] <= 8000
assert 5000 <= office["word_count"] <= 8000

sys.path.insert(0, str(RELEASE))
from experiments.lab02.eval.artifacts import execution_manifest
current = execution_manifest(RELEASE)
assert current["source_sha256"] == snapshot["source_sha256"]
assert current["snapshot_sha256"] == snapshot["raw_data_sha256"]
for name, expected in snapshot["runtime_script_sha256"].items():
    assert hashlib.sha256((RELEASE / name).read_bytes()).hexdigest() == expected

table_files = {
    "forward": "benchmarks/forward_summary.csv", "slicing": "benchmarks/goal_slicing.csv",
    "resolution": "benchmarks/resolution.csv", "indicator": "experiments/summary.csv",
    "sensitivity": "experiments/sensitivity.csv", "ablation": "experiments/ablations.csv",
    "conflict": "experiments/conflict_summary.csv", "graph": "end_to_end/summary.csv",
    "equity": "governance/equity_summary.csv",
}
for key, relative in table_files.items():
    with (RELEASE / "outputs/lab02" / relative).open(encoding="utf-8-sig", newline="") as file:
        assert list(csv.DictReader(file)) == snapshot["tables"][key], key

markdown = (WORK / "report-final.md").read_text(encoding="utf-8")
body, references = markdown.split("# 参考文献", 1)
assert list(dict.fromkeys(re.findall(r"\[(\d+)\]", body))) == list(map(str, range(1, 9)))
assert not re.search(r"@@[A-Z_]+@@", markdown)
for stale in ("362.81", "57.75", "6.28倍", "200项Python测试", "根目录的README"):
    assert stale not in markdown, stale

document = Document(WORK / (STEM + ".docx"))
headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
for required in ("摘要", "1 引言", "2 相关理论基础", "3 场景与模型建模", "4 算法设计与工程实现",
                 "5 实验设计与评价方法", "6 实验结果与分析", "7 创新点总结", "8 不足、展望与结论",
                 "参考文献", "附录A 小组分工", "附录B 环境配置、运行与结果解读"):
    assert required in headings, required
assert len(document.tables) == 7 and len(document.inline_shapes) == 4
assert len(document.element.xpath(".//m:oMath")) >= 4
with zipfile.ZipFile(WORK / (STEM + ".docx")) as archive:
    assert archive.testzip() is None
pdf = fitz.open(WORK / (STEM + ".pdf"))
assert len(pdf) == office["pages"] == 17
layout = json.loads((WORK / "layout-review.json").read_text(encoding="utf-8"))
assert all(not page["outside_page"] for page in layout)
assert all(page["text_characters"] > 100 for page in layout)

example = re.search(r"```python\n(.*?)\n```", markdown, flags=re.S).group(1)
namespace = {}
exec(compile(example, "report_example", "exec"), namespace)
assert namespace["result"]["company_id"] == "C3" and namespace["result"]["proofs"]

protected = json.loads((WORK / "protected-pptx.json").read_text(encoding="utf-8"))
assert all(hashlib.sha256(Path(path).read_bytes()).hexdigest() == digest for path, digest in protected.items())
DESTINATION.mkdir(exist_ok=True)
published = []
for extension in (".docx", ".pdf"):
    source = WORK / (STEM + extension)
    target = DESTINATION / source.name
    if target.exists():
        backup = WORK / ("before-publish-" + time.strftime("%Y%m%d-%H%M%S") + extension)
        shutil.copy2(target, backup)
    shutil.copy2(source, target)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    assert hashlib.sha256(target.read_bytes()).hexdigest() == digest
    published.append({"path": str(target), "bytes": target.stat().st_size, "sha256": digest})
report = {"status": "complete", "report_only": True, **office,
          "source_files_checked": len(current["source_sha256"]), "experiment_tables_checked": len(table_files),
          "data_unchanged_during_writing": True, "figures": 4, "tables_including_cover_and_team": 7,
          "references": 8, "example_code_runs": True, "layout_review_passed": True,
          "editable_equations": len(document.element.xpath(".//m:oMath")), "published": published}
(WORK / "completion.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
