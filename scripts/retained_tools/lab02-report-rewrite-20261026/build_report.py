"""Build only the rewritten report; the submitted code and presentation are read-only inputs."""
import importlib.util
import json
from pathlib import Path
import re
import shutil
import statistics
import subprocess

from docx import Document
from docx.shared import Cm, Pt
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter
import numpy as np

WORK = Path(__file__).resolve().parent
PROJECT = WORK.parents[2]
SNAPSHOT = json.loads((WORK / "evidence.json").read_text(encoding="utf-8"))
DATA = SNAPSHOT["tables"]
BRIEF = json.loads((WORK / "brief.json").read_text(encoding="utf-8"))
STEM = "第31组_人工智能推理实验报告"
FIGURES = WORK / "figures"
FIGURES.mkdir(exist_ok=True)
plt.rcParams.update({"font.family": "Microsoft YaHei", "font.size": 10.5, "axes.titlesize": 11,
                     "axes.labelsize": 10.5, "legend.fontsize": 10, "axes.unicode_minus": False,
                     "axes.spines.top": False, "axes.spines.right": False,
                     "figure.facecolor": "white", "savefig.facecolor": "white"})
BLUE, RUST, GRAY = "#356D86", "#B25835", "#87918F"


def row(table, **conditions):
    return next(r for r in DATA[table] if all(r[k] == v for k, v in conditions.items()))


def save_figure(fig, filename):
    fig.savefig(FIGURES / filename, dpi=280, bbox_inches="tight", pad_inches=.12)
    plt.close(fig)


def make_figures():
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.05), layout="constrained")
    methods = [("naive_scan", "全扫描", GRAY), ("indexed_naive", "索引朴素", BLUE), ("semi_naive", "半朴素", RUST)]
    for model, label, color in methods:
        values = [row("forward", groups=str(n), model=model) for n in (1, 4, 8, 16)]
        axes[0].errorbar((1, 4, 8, 16), [float(v["elapsed_ms_mean"]) for v in values],
                         yerr=[float(v["elapsed_ms_sd"]) for v in values], marker="o", ms=4,
                         lw=1.7, capsize=2, label=label, color=color)
    axes[0].set(yscale="log", xlabel="结构组数", ylabel="耗时 / ms（对数）", xticks=[1, 4, 8, 16], title="(a) 规模与运行时间")
    axes[0].legend(frameon=False, fontsize=9)
    counts = [float(row("forward", groups="16", model=model)["candidate_checks_mean"]) for model, _, _ in methods]
    axes[1].bar(range(3), counts, color=[m[2] for m in methods], width=.60)
    axes[1].set(yscale="log", ylim=(1000, 7e8), xticks=range(3), xticklabels=[m[1] for m in methods], ylabel="候选检查次数（对数）", title="(b) 16组结构的匹配工作量")
    for x, value in enumerate(counts):
        axes[1].text(x, value * 1.4, f"{int(value):,}", ha="center", va="bottom", fontsize=9)
    for ax in axes:
        ax.grid(axis="y", alpha=.18)
        ax.set_axisbelow(True)
    save_figure(fig, "01_deterministic.png")

    fig, ax = plt.subplots(figsize=(7.2, 3.0), layout="constrained")
    scenarios = ["clean", "noise_010", "noise_020", "noise_030", "missing_010", "missing_030", "mixed", "severe"]
    labels = ["清洁", "噪声\n0.10", "噪声\n0.20", "噪声\n0.30", "缺测\n10%", "缺测\n30%", "混合", "重度"]
    for model, label, color, marker in [("bayes_learned", "学习Bayes", BLUE, "o"), ("fuzzy_learned", "拟合Mamdani", RUST, "s")]:
        rows = [row("indicator", scenario=s, model=model) for s in scenarios]
        ax.errorbar(range(8), [float(r["f1_mean"]) for r in rows], yerr=[float(r["f1_sd"]) for r in rows],
                    label=label, color=color, marker=marker, ms=4, lw=1.8, capsize=3)
    ax.set(xticks=range(8), xticklabels=labels, ylim=(.5, .95), ylabel="F1", xlabel="观测条件")
    ax.legend(frameon=False, ncol=2, loc="upper right")
    ax.grid(axis="y", alpha=.2)
    save_figure(fig, "02_robustness.png")

    fig, axes = plt.subplots(1, 3, figsize=(7.2, 2.65), layout="constrained", sharey=True)
    for ax, model, parameter, label, color in zip(axes, ("bayes", "bayes", "fuzzy"),
            ("prior", "cpt_positive_shift", "high_start"), ("(a) Bayes先验", "(b) 正类CPT偏移", "(c) 高隶属起点"), (BLUE, BLUE, RUST)):
        rows = [r for r in SNAPSHOT["sensitivity_summary"] if r["model"] == model and r["parameter"] == parameter]
        ax.errorbar([r["value"] for r in rows], [r["mean"] for r in rows], yerr=[r["sd"] for r in rows],
                    marker="o", color=color, lw=1.5, ms=4, capsize=2)
        ax.set(title=label, xlabel="参数取值", ylim=(.60, .94))
        ax.grid(axis="y", alpha=.2)
    axes[0].set_ylabel("F1")
    save_figure(fig, "03_sensitivity.png")

    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.15), layout="constrained")
    scenarios = ["clean", "false_loop_020", "severe"]
    policies = [("hybrid_adaptive_unlocked", "不锁定", GRAY), ("hybrid_adaptive_locked", "硬锁定", BLUE), ("hybrid_adaptive_calibrated", "校准复核", RUST)]
    for ax, metric, title in zip(axes, ("f1", "false_positive_rate"), ("(a) 最终行动的F1", "(b) 最终行动的误报率")):
        for index, (model, label, color) in enumerate(policies):
            rows = [row("graph", scenario=s, model=model) for s in scenarios]
            ax.bar(np.arange(3) + (index - 1) * .24, [float(r[metric + "_mean"]) for r in rows], width=.23,
                   yerr=[float(r[metric + "_sd"]) for r in rows], capsize=2, color=color, label=label)
        ax.set(xticks=range(3), xticklabels=["清洁", "错误闭环", "重度"], title=title, ylim=(0, 1 if metric == "f1" else .65))
        ax.grid(axis="y", alpha=.18)
        ax.set_axisbelow(True)
    axes[1].yaxis.set_major_formatter(PercentFormatter(1))
    axes[0].legend(frameon=False, ncol=3, loc="upper left", fontsize=8.5)
    save_figure(fig, "04_policy.png")


def table(caption, headers, rows):
    return caption + "\n\n| " + " | ".join(headers) + " |\n|" + "|".join("---" for _ in headers) + "|\n" + "\n".join("| " + " | ".join(map(str, r)) + " |" for r in rows)


def replacements():
    fw = []
    for model, label in [("naive_scan", "全扫描"), ("indexed_naive", "索引朴素"), ("semi_naive", "半朴素")]:
        r = row("forward", groups="16", model=model)
        fw.append([label, f"{float(r['elapsed_ms_mean']):.2f}±{float(r['elapsed_ms_sd']):.2f}",
                   f"{int(float(r['candidate_checks_mean'])):,}", "20", "3031"])
    indicators = []
    for model, label in [("bayes_manual", "手工Bayes"), ("bayes_learned", "学习Bayes"), ("fuzzy", "手工Mamdani"),
                         ("fuzzy_learned", "拟合Mamdani"), ("hybrid_fixed", "固定混合"), ("hybrid_adaptive", "自适应混合")]:
        c, s = row("indicator", scenario="clean", model=model), row("indicator", scenario="severe", model=model)
        indicators.append([label, f"{float(c['f1_mean']):.4f}", f"{float(c['brier_mean']):.4f}",
                           f"{float(s['f1_mean']):.4f}", f"{float(c['inference_us_mean']):.2f}"])
    policies = []
    for scenario, label in [("clean", "清洁"), ("false_loop_020", "错误闭环"), ("severe", "重度")]:
        for model, name in [("hybrid_adaptive_unlocked", "不锁定"), ("hybrid_adaptive_locked", "硬锁定"), ("hybrid_adaptive_calibrated", "校准复核")]:
            r = row("graph", scenario=scenario, model=model)
            policies.append([label, name, f"{float(r['recall_mean'])*100:.2f}", f"{float(r['false_positive_rate_mean'])*100:.2f}", f"{float(r['f1_mean']):.4f}"])
    comparison = [
        ["建模难度", "明确谓词、规则、守卫与业务边界。", "Bayes需依赖结构和CPT；模糊方法需集合与规则。"],
        ["推理效率", "取决于连接、递归和规模；可用索引、增量和切片优化。", "星形Bayes较轻；模糊聚合随规则和网格开销增加。"],
        ["容错性", "缺事实可能漏推，错误事实可能传播；已证正例仍保留。", "可处理部分观测；参数偏差和分布变化仍使结果退化。"],
        ["结果可信度", "证明可追溯，但正确性有赖于事实与规则的有效性。", "概率与隶属度需区别；校准和数据质量决定解释范围。"],
        ["落地限制", "给定规则不等于完整业务规范，不能自动核验输入。", "需外部标签、来源质量和决策代价，融合不保证全面增益。"],
    ]
    features = [
        ["不透明度", "登记披露的教学测度", "登记资料缺失时为未知"],
        ["关联交易", "已识别关联交易额／注册资本", "持股或交易资料不完整时为未知"],
        ["担保负荷", "对外担保额／注册资本", "担保资料不完整时为未知"],
        ["回路线索", "持股或担保可达闭包回到自身", "保留已证正例；缺测且未证时为未知"],
    ]
    return {
        "FEATURE_TABLE": table("表1　共同指标及缺测处理", ["指标", "建模定义", "缺测语义"], features),
        "FORWARD_TABLE": table("表2　16组结构的完整闭包基准（三次重复，均值±标准差）", ["方法", "平均耗时／ms", "候选检查次数", "迭代轮数", "事实数"], fw),
        "INDICATOR_TABLE": table("表3　指标层方法比较（三个种子均值）", ["方法", "清洁F1", "清洁Brier", "重度F1", "清洁耗时／μs"], indicators),
        "POLICY_TABLE": table("表4　共用自适应评分的图级策略比较（三个种子均值）", ["条件", "行动策略", "召回率／%", "误报率／%", "F1"], policies),
        "COMPARISON_TABLE": table("表5　确定性与不确定性推理的五维比较", ["维度", "确定性推理", "不确定性推理"], comparison),
        "TEAM_TABLE": table("小组责任分工安排", ["姓名", "学号", "职能", "责任范围"],
                            [[m["name"], m["id"], m["role"], m["responsibility"]] for m in BRIEF["members"]]),
        "SPEEDUP": f"{SNAPSHOT['speedup']:.2f}",
        "INDEX_MS": f"{float(row('forward', groups='16', model='indexed_naive')['elapsed_ms_mean']):.2f}",
        "SEMI_MS": f"{float(row('forward', groups='16', model='semi_naive')['elapsed_ms_mean']):.2f}",
        "BAYES_US": f"{float(row('indicator', scenario='clean', model='bayes_learned')['inference_us_mean']):.2f}",
        "FUZZY_US": f"{float(row('indicator', scenario='clean', model='fuzzy_learned')['inference_us_mean']):.2f}",
        "RESOLUTION_BASE": f"{statistics.mean(float(r['elapsed_ms']) for r in DATA['resolution'] if r['model']=='fifo_baseline'):.2f}",
        "RESOLUTION_OPT": f"{statistics.mean(float(r['elapsed_ms']) for r in DATA['resolution'] if r['model']=='optimized'):.2f}",
    }


def main():
    make_figures()
    source = (WORK / "report-source.md").read_text(encoding="utf-8")
    for key, value in replacements().items():
        source = source.replace("@@" + key + "@@", value)
    assert not re.search(r"@@[A-Z_]+@@", source)
    markdown = WORK / "report-final.md"
    markdown.write_text(source, encoding="utf-8")

    spec = importlib.util.spec_from_file_location("report_format", PROJECT / "outputs/qa/lab02-materials/build_documents.py")
    formatter = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(formatter)
    formatter.WORK = WORK
    formatter.BRIEF = BRIEF
    reference = formatter.reference_document()
    intermediate = WORK / "report-pandoc.docx"
    command = [shutil.which("pandoc"), str(markdown), "-o", str(intermediate),
               "--reference-doc=" + str(reference), "--resource-path=" + str(WORK),
               "-M", "lang=zh-CN", "--highlight-style=tango"]
    subprocess.run(command, check=True, capture_output=True, text=True, encoding="utf-8")
    document = Document(intermediate)
    formatter.cover(document)
    formatter.format_document(document, "report")
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("关键词："):
            paragraph.paragraph_format.first_line_indent = Pt(0)
        if paragraph.text.startswith("附录B"):
            paragraph.paragraph_format.page_break_before = True
    paragraphs = document.paragraphs
    for previous, following in zip(paragraphs, paragraphs[1:]):
        if following._p.xpath(".//m:oMathPara"):
            previous.paragraph_format.keep_with_next = True
    target = WORK / (STEM + ".docx")
    document.save(target)
    all_text = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
    result = {"file": str(target), "tables": len(document.tables), "figures": len(document.inline_shapes),
              "editable_equations": len(document.element.xpath(".//m:oMath")),
              "cjk_characters": len(re.findall(r"[\u4e00-\u9fff]", all_text)),
              "nonspace_characters": len(re.sub(r"\s", "", all_text))}
    (WORK / "build.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
