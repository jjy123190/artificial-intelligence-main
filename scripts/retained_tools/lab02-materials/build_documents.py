"""Build editable academic documents using the supplied reference layout."""
from __future__ import annotations
import copy
import argparse
import json
import re
import shutil
import subprocess
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

WORK = Path(__file__).resolve().parent
ROOT = WORK.parents[2]
OUT = ROOT / "deliverables/lab02"
BRIEF = json.loads((WORK / "brief.json").read_text(encoding="utf-8"))
DATA = json.loads((WORK / "data-snapshot.json").read_text(encoding="utf-8"))["tables"]
GROUP = f"第{BRIEF['group']}组"
NAMES = {"report": GROUP+"_人工智能推理实验报告", "innovation": GROUP+"_创新点说明", "talk": GROUP+"_答辩讲稿与问答"}


def md_table(caption, headers, rows):
    return caption + "\n\n| " + " | ".join(headers) + " |\n|" + "|".join("---" for _ in headers) + "|\n" + "\n".join("| " + " | ".join(str(x) for x in row) + " |" for row in rows)


def row(family, scenario, model):
    return next(r for r in DATA[family] if r["scenario"] == scenario and r["model"] == model)


def substitutions():
    features = [["不透明度", "登记与披露信息的教学测度", "登记测度缺失时为未知"],
                ["关联交易比例", "已识别关联交易额／注册资本，截断到[0,1]", "持股或交易资料不完整时为未知"],
                ["担保负荷", "对外担保额／注册资本，截断到[0,1]", "担保资料不完整时为未知"],
                ["回路线索", "持股或担保可达闭包是否回到自身", "已证正例保留；缺测且未证时为未知"]]
    fw = []
    for model,label in [("naive_scan","全扫描"),("indexed_naive","索引朴素"),("semi_naive","半朴素")]:
        r=next(r for r in DATA["forward"] if r["groups"]=="16" and r["model"]==model)
        fw.append([label,f"{float(r['elapsed_ms_mean']):.2f}±{float(r['elapsed_ms_sd']):.2f}",
                   f"{int(float(r['candidate_checks_mean'])):,}",int(float(r['cycles_mean'])),int(float(r['facts_mean']))])
    indicators=[]
    for model,label in [("bayes_manual","Bayes手工"),("bayes_learned","Bayes学习"),("fuzzy","Mamdani手工"),
                         ("fuzzy_learned","Mamdani拟合"),("hybrid_fixed","固定混合")]:
        clean=row("indicator","clean",model)
        indicators.append([label,f"{float(clean['f1_mean']):.4f}",f"{float(clean['brier_mean']):.4f}",
                           f"{float(row('indicator','mixed',model)['f1_mean']):.4f}",f"{float(row('indicator','severe',model)['f1_mean']):.4f}"])
    policies=[]
    for scenario,label in [("clean","清洁"),("false_loop_020","错误闭环"),("severe","重度")]:
        for model,name in [("hybrid_adaptive_unlocked","不锁定"),("hybrid_adaptive_locked","硬锁定"),("hybrid_adaptive_calibrated","校准复核")]:
            r=row("graph",scenario,model)
            policies.append([label,name,f"{float(r['recall_mean'])*100:.2f}",f"{float(r['false_positive_rate_mean'])*100:.2f}",f"{float(r['f1_mean']):.4f}"])
    compare=[["建模难度","需明确事实、谓词、规则与数值守卫，业务边界清楚。","Bayes需图结构和CPT；模糊方法需隶属函数与规则；混合还需行动策略。"],
             ["推理效率","受前件连接与递归规模影响；索引、增量与切片可减少工作量。","本例星形Bayes较轻；Mamdani含网格聚合；一般网络复杂度受树宽影响。"],
             ["容错能力","已有正向证据可保留，但缺测会漏推，错误事实会传播。","可处理部分观测并连续评分；参数偏差与错误结构仍可能增加误报。"],
             ["结果可信度","结论可重放，但逻辑可靠性以输入事实和规则有效为前提。","后验、隶属度和最终行动需分开解释；评分未必经过概率校准。"],
             ["落地限制","不能独立表达测量概率，教学规则未覆盖全部真实业务。","需要可信参数、数据分布与代价设定；融合不保证普遍优于单模型。"]]
    team=[[m["name"],m["id"],m["role"],m["responsibility"]] for m in BRIEF["members"]]
    return {"FEATURE_TABLE":md_table("表1　共同指标及缺测语义",["指标","建模定义","缺测处理"],features),
            "FORWARD_TABLE":md_table("表2　16组结构的完整闭包基准（三次重复，均值±标准差）",["方法","平均耗时／ms","候选检查次数","迭代轮数","事实数"],fw),
            "INDICATOR_TABLE":md_table("表3　指标层方法比较（三个种子均值）",["方法","清洁F1","清洁Brier","混合F1","重度F1"],indicators),
            "POLICY_TABLE":md_table("表4　自适应评分模型的图级策略比较（三个种子均值）",["条件","行动策略","召回率／%","误报率／%","F1"],policies),
            "COMPARISON_TABLE":md_table("表5　同场景确定性与不确定性方法的五维比较",["维度","确定性规则推理","不确定性与混合推理"],compare),
            "TEAM_TABLE":md_table("小组责任分工安排",["姓名","学号","职能","责任范围"],team)}


def font_style(style, size=12, bold=False, east="宋体", latin="Times New Roman"):
    style.font.name=latin;style.font.size=Pt(size);style.font.bold=bold
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),east)


def reference_document():
    template=next(ROOT.glob("*.docx"))
    d=Document(template)
    for child in list(d.element.body):
        if child.tag != qn("w:sectPr"):d.element.body.remove(child)
    for name in ["Body Text","First Paragraph","Heading 1","Heading 2","Heading 3","Caption","Image Caption","Table Caption","Source Code"]:
        if name not in d.styles:
            style=d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            style.base_style=d.styles["Normal"]
    for name in ["Normal","Body Text","First Paragraph"]:
        if name not in d.styles:continue
        s=d.styles[name];font_style(s)
        s.paragraph_format.first_line_indent=Pt(24)
        s.paragraph_format.line_spacing=1.5
        s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(4)
        s.paragraph_format.widow_control=True
    for name,size in [("Heading 1",15),("Heading 2",12),("Heading 3",11)]:
        s=d.styles[name];font_style(s,size,True,"黑体")
        s.paragraph_format.first_line_indent=Pt(0)
        s.paragraph_format.space_before=Pt(14);s.paragraph_format.space_after=Pt(7)
        s.paragraph_format.keep_with_next=True
    for name in ["Caption","Image Caption","Table Caption"]:
        if name in d.styles:
            s=d.styles[name];font_style(s,10.5)
            s.paragraph_format.first_line_indent=Pt(0)
    ref=WORK/"reference-styled.docx";d.save(ref)
    return ref


def set_run_font(run,size=12,bold=None,east="宋体",latin="Times New Roman"):
    run.font.name=latin;run.font.size=Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),east)
    if bold is not None:run.bold=bold


def cover(d):
    original=list(d.element.body)[:-1]
    def line(text,size=12,bold=False,before=0,after=8):
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.25
        p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
        set_run_font(p.add_run(text),size,bold,"黑体" if bold else "宋体")
        return p
    line(BRIEF["school"],24,True,28,14)
    line(BRIEF["college"],13,False,0,28)
    line("《人工智能导论》\n推理技术分组创新作业",19,True,0,25)
    line("面向企业股权风险评估的\n确定性与不确定性混合推理",21,True,0,26)
    line(BRIEF["class"]+"　　第"+BRIEF["group"]+"组",13,False,0,12)
    t=d.add_table(rows=1,cols=3)
    for cell,text in zip(t.rows[0].cells,["姓名","学号","职能"]):cell.text=text
    for m in BRIEF["members"]:
        for cell,text in zip(t.add_row().cells,[m["name"],m["id"],m["role"]]):cell.text=text
    line("组长：梁镇贤　　课堂汇报：刘亦达",12,False,12,12)
    line(BRIEF["cover_date"],12,False,12,0)
    d.add_section(WD_SECTION_START.NEW_PAGE)
    additions=[e for e in list(d.element.body)[:-1] if e not in original]
    for e in additions:d.element.body.remove(e)
    for index,e in enumerate(additions):d.element.body.insert(index,e)


def cell_shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr();shd=OxmlElement("w:shd");shd.set(qn("w:fill"),fill);pr.append(shd)


def format_document(d,key):
    for p in d.paragraphs:
        name=p.style.name
        if name.startswith("Heading"):
            p.paragraph_format.keep_with_next=True
            if p.text.startswith(("参考文献","附录A")):p.paragraph_format.page_break_before=True
        elif name in ("Source Code","Verbatim Char"):
            p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.05
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.keep_together=True
            for r in p.runs:set_run_font(r,8.5,False,"微软雅黑","Consolas")
        elif name in ("Body Text","First Paragraph","Normal") and p.alignment!=WD_ALIGN_PARAGRAPH.CENTER:
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent=Pt(24)
        if re.match(r"^(表\d|小组责任分工)",p.text):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.keep_with_next=True
            for r in p.runs:set_run_font(r,10.5,False)
        if p._p.xpath(".//w:drawing"):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.keep_with_next=True
        if name in ("Caption","Image Caption"):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
            for r in p.runs:set_run_font(r,10.5,False)
        if re.match(r"^\[\d\]",p.text):
            p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.15
            p.paragraph_format.space_after=Pt(7)
            for r in p.runs:set_run_font(r,10.5)
    for table in d.tables:
        count=len(table.columns);heads=[c.text for c in table.rows[0].cells]
        if count==5:
            widths=[2.3,4.0,3.05,2.65,2.65] if "平均耗时／ms" in heads else ([2.2,4.05,2.8,2.8,2.8] if heads[0]=="条件" else [4.25,2.6,2.6,2.6,2.6])
        elif count==4:widths=[1.65,3.0,2.75,7.25]
        elif heads[0]=="姓名":widths=[2.2,5.0,7.45]
        else:widths=[2.4,6.125,6.125]
        table.autofit=False
        for i,w in enumerate(widths):table.columns[i].width=Cm(w)
        for ri,r in enumerate(table.rows):
            pr=r._tr.get_or_add_trPr();cant=OxmlElement("w:cantSplit");pr.append(cant)
            if ri==0:pr.append(OxmlElement("w:tblHeader"))
            for i,c in enumerate(r.cells):
                c.width=Cm(widths[i])
                if ri==0:cell_shade(c,"EEE8E0")
                for p in c.paragraphs:
                    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
                    p.paragraph_format.line_spacing=1.12
                    p.paragraph_format.keep_with_next=ri<len(table.rows)-1
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if count==4 or (count==3 and heads[0]!="姓名") else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:set_run_font(run,10 if count>=4 else 10.5,ri==0)
    sections=d.sections
    for i,s in enumerate(sections):
        if key=="report" and i==0:continue
        s.header.is_linked_to_previous=False;s.footer.is_linked_to_previous=False
        p=s.header.paragraphs[0];p.text="华南农业大学｜人工智能导论｜"+GROUP
        p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:set_run_font(r,9);r.font.color.rgb=RGBColor.from_string("737373")
        p=s.footer.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run();fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");r._r.addnext(fld)
        num=OxmlElement("w:pgNumType");num.set(qn("w:start"),"1");s._sectPr.append(num)
    d.core_properties.title=BRIEF["title"] if key=="report" else NAMES[key]
    d.core_properties.author="；".join(m["name"] for m in BRIEF["members"])
    d.core_properties.subject="人工智能导论推理技术课程作业，"+GROUP
    d.core_properties.keywords="确定性推理,贝叶斯推理,模糊推理,混合推理,股权风险"


def main():
    parser=argparse.ArgumentParser();parser.add_argument("--only",choices=["report","innovation","talk"]);args=parser.parse_args()
    OUT.mkdir(parents=True,exist_ok=True)
    ref=reference_document();replace=substitutions();audit_path=WORK/"documents-built.json"
    audit=json.loads(audit_path.read_text(encoding="utf-8")) if args.only and audit_path.exists() else []
    for key in ([args.only] if args.only else ["report","innovation","talk"]):
        source=(WORK/(key+"-source.md")).read_text(encoding="utf-8")
        for name,value in replace.items():source=source.replace("{{"+name+"}}",value)
        assert not re.search(r"\{\{[A-Z_]+\}\}",source)
        assert "刘奕达" not in source and "梁镇闲" not in source
        if key!="report":source=source.replace("表4　自适应", "自适应")
        final_md=WORK/(key+"-final.md");final_md.write_text(source,encoding="utf-8")
        intermediate=WORK/(key+"-pandoc.docx")
        cmd=[shutil.which("pandoc"),str(final_md),"-o",str(intermediate),"--reference-doc="+str(ref),
             "--resource-path="+str(WORK),"--shift-heading-level-by=-1","-M","lang=zh-CN","--highlight-style=tango"]
        completed=subprocess.run(cmd,capture_output=True,text=True,encoding="utf-8")
        if completed.returncode:raise RuntimeError(completed.stderr)
        d=Document(intermediate)
        if key=="report":cover(d)
        format_document(d,key)
        target=OUT/(NAMES[key]+".docx");d.save(target)
        all_text="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
        rec={"kind":key,"file":target.name,"cjk_chars":len(re.findall(r"[\u4e00-\u9fff]",all_text)),
             "nonspace_chars":len(re.sub(r"\s","",all_text)),"tables":len(d.tables),"figures":len(d.inline_shapes),
             "equations":len(d.element.xpath(".//m:oMath")),"sections":len(d.sections)}
        audit=[r for r in audit if r["kind"]!=key]+[rec];print(json.dumps(rec,ensure_ascii=False),flush=True)
    (WORK/"documents-built.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding="utf-8")


if __name__=="__main__":main()
