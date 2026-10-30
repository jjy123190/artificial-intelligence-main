"""Native editable PowerPoint: text, shapes, connectors, tables and data charts."""
from pathlib import Path
import json
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_TICK_MARK
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement

WORK=Path(__file__).resolve().parent
ROOT=WORK.parents[2]
OUT=ROOT/"deliverables/lab02"
B=json.loads((WORK/"brief.json").read_text(encoding="utf-8"))
GROUP=f"第{B['group']}组"
D=json.loads((WORK/"data-snapshot.json").read_text(encoding="utf-8"))["tables"]
BG="F8F5F0";INK="292724";MUTED="71685F";ACCENT="9A4728";LINE="D9CFC3";BLUE="426D80";GREEN="54785F";GRAY="939DA0";PALE="EFE6DB"
FONT="Microsoft YaHei"


def rgb(code):return RGBColor.from_string(code)


def text(s,x,y,w,h,value,size=20,color=INK,bold=False,font=FONT,align=PP_ALIGN.LEFT):
    shape=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(5);p.line_spacing=1.12
        r=p.add_run();r.text=line;r.font.name=font;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
        pr=r._r.get_or_add_rPr();pr.set("lang","zh-CN")
        ea=OxmlElement("a:ea");ea.set("typeface",font);pr.append(ea)
    return shape


def rectangle(s,x,y,w,h,fill=PALE,line=None,rounded=False):
    shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                             Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill)
    if line:shape.line.color.rgb=rgb(line);shape.line.width=Pt(.8)
    else:shape.line.fill.background()
    return shape


def connector(s,x1,y1,x2,y2,color=ACCENT,width=1.6,arrow=False):
    shape=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb=rgb(color);shape.line.width=Pt(width)
    if arrow:
        end=OxmlElement("a:tailEnd");end.set("type","triangle");shape._element.spPr.get_or_add_ln().append(end)
    return shape


def standard(prs,number,title,kicker="实验二 · 推理技术",backup=False):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
    text(s,.65,.38,10,.27,kicker,11,ACCENT,True)
    text(s,.65,.86,12.05,.76,title,31,INK,True)
    connector(s,.65,6.95,12.68,6.95,LINE,.7)
    text(s,.65,7.08,10,.20,"华南农业大学｜人工智能导论｜"+GROUP,9.5,MUTED)
    text(s,11.5,7.05,1.18,.28,f"备用 {number-6}/2" if backup else f"{number:02d} / 06",10,MUTED,align=PP_ALIGN.RIGHT)
    return s


def note(s,number):
    src=(WORK/"talk-source.md").read_text(encoding="utf-8")
    match=re.search(rf"#### 第{number}页[^\n]*\n\n(.*?)(?=\n####|\n###|\Z)",src,re.S)
    body=match.group(1).strip() if match else "答辩备用页；按老师问题展开，不计入主讲时间。"
    s.notes_slide.notes_text_frame.text="汇报人：刘亦达\n"+body


def setup_chart(chart):
    chart.has_legend=False;chart.chart_style=10
    chart.font.name=FONT;chart.font.size=Pt(13)
    chart.category_axis.tick_labels.font.name=FONT;chart.category_axis.tick_labels.font.size=Pt(16)
    chart.value_axis.tick_labels.font.name=FONT;chart.value_axis.tick_labels.font.size=Pt(11)
    chart.value_axis.has_major_gridlines=True
    chart.value_axis.major_gridlines.format.line.color.rgb=rgb(LINE)
    chart.category_axis.format.line.color.rgb=rgb(LINE)
    chart.value_axis.format.line.color.rgb=rgb(LINE)
    chart.category_axis.major_tick_mark=XL_TICK_MARK.NONE
    chart.value_axis.major_tick_mark=XL_TICK_MARK.NONE
    plot=chart.plots[0];plot.gap_width=85;plot.has_data_labels=True
    plot.data_labels.font.name=FONT;plot.data_labels.font.size=Pt(14)
    plot.data_labels.position=XL_DATA_LABEL_POSITION.OUTSIDE_END
    for point,color in zip(chart.series[0].points,[GRAY,BLUE,ACCENT]):
        point.format.fill.solid();point.format.fill.fore_color.rgb=rgb(color);point.format.line.fill.background()


def main():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    s=standard(prs,1,"面向企业股权风险评估的\n确定性与不确定性混合推理")
    # The title gets a taller native text box on the cover.
    title=s.shapes[1];title.height=Inches(1.55)
    text(s,.68,2.70,8.7,.8,"证明结构 · 量化不确定性 · 比较行动代价",24,ACCENT,True)
    text(s,.70,3.68,8.5,.70,"数学与信息学院、软件学院\n软件工程5班",17,MUTED)
    text(s,.70,4.65,9,.38,"　".join(m["name"] for m in B["members"]),17,INK)
    text(s,.70,5.24,8.7,.34,"组长：梁镇贤　　汇报：刘亦达",16,INK)
    text(s,.70,5.88,8.0,.3,B["cover_date"],14,MUTED)
    text(s,10.0,2.65,2.55,2.6,B["group"],132,"D5C3B2",True,align=PP_ALIGN.CENTER)
    text(s,10.0,5.35,2.6,.5,GROUP,17,MUTED,align=PP_ALIGN.CENTER)
    note(s,1)

    s=standard(prs,2,"控制关系与经济权益，分别回答")
    text(s,.68,1.76,11.5,.55,"同源案例 → 27条安全正Horn规则 → 可重放的推理路径",19,MUTED)
    xs=[.9,3.15,5.4,7.65];y=3.0;diam=.88
    for x,name in zip(xs,["P1","C1","C2","C3"]):
        node=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(diam),Inches(diam))
        node.fill.solid();node.fill.fore_color.rgb=rgb(ACCENT if name=="C3" else "E6D9CA");node.line.fill.background()
        text(s,x,y+.20,diam,.42,name,22,"FFFFFF" if name=="C3" else INK,True,align=PP_ALIGN.CENTER)
    for i,label in enumerate(["80%","65%","35%"]):
        connector(s,xs[i]+diam,y+.44,xs[i+1]-.05,y+.44,ACCENT,1.8,True)
        text(s,xs[i]+.95,y-.48,1.0,.36,label,18,ACCENT,True,align=PP_ALIGN.CENTER)
    connector(s,3.59,3.91,3.59,4.65,BLUE,1.5)
    connector(s,3.59,4.65,8.09,4.65,BLUE,1.5)
    connector(s,8.09,4.65,8.09,3.91,BLUE,1.5,True)
    text(s,5.03,4.8,1.6,.40,"直接持股20%",16,BLUE)
    text(s,9.5,2.52,2.9,.74,"55%",42,ACCENT,True)
    text(s,9.5,3.30,2.9,.8,"C1可支配份额\n35% + 20%",17,MUTED)
    text(s,9.5,4.35,2.9,.75,"34.2%",38,BLUE,True)
    text(s,9.5,5.12,2.9,.50,"P1简单路径权益",17,MUTED)
    text(s,.74,5.8,8.1,.54,"0.8 × 0.2 + 0.8 × 0.65 × 0.35 = 0.342",20,INK)
    text(s,.74,6.43,11.5,.28,"控制候选取决于给定规则；逻辑可靠性不等于输入资料真实性。",12,MUTED)
    note(s,2)

    s=standard(prs,3,"同一组指标，两种不确定性表达")
    rectangle(s,.68,1.83,11.95,.57,"EDE6DD")
    text(s,.91,1.97,11.45,.28,"不透明度　／　关联交易比例　／　担保负荷　／　回路线索",16,MUTED)
    connector(s,6.58,2.76,6.58,5.91,LINE,1)
    text(s,.84,2.88,5.35,.55,"Bayes 概率推理",26,BLUE,True)
    text(s,.85,3.60,5.35,.60,"先验 × 似然 → 后验",23,INK)
    text(s,.87,4.35,5.35,1.25,"训练集估计先验与CPT\n缺测指标从证据中省略\n清洁指标实验 F1：0.8575",19,MUTED)
    text(s,7.01,2.88,5.3,.55,"Mamdani 模糊推理",26,ACCENT,True)
    text(s,7.02,3.60,5.30,.60,"隶属度 → 规则 → 质心",23,INK)
    text(s,7.04,4.35,5.30,1.25,"训练集拟合15组预设候选\n缺测评分向中性值收缩\n清洁指标实验 F1：0.7053",19,MUTED)
    text(s,.87,6.35,11.60,.38,"概率与模糊评分分别解释；模型选择仅使用训练集和验证集。",16,INK)
    note(s,3)

    s=standard(prs,4,"相同闭包，减少重复匹配")
    text(s,.72,1.77,11.90,.50,"16组股权结构｜3031条事实｜20轮迭代｜三次重复实测",18,MUTED)
    rows=[next(r for r in D["forward"] if r["groups"]=="16" and r["model"]==m) for m in ["naive_scan","indexed_naive","semi_naive"]]
    data=CategoryChartData();data.categories=["全扫描","索引朴素","半朴素"]
    data.add_series("平均耗时/ms",[float(r["elapsed_ms_mean"]) for r in rows])
    chart=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(.72),Inches(2.54),Inches(7.62),Inches(3.65),data).chart
    setup_chart(chart);chart.value_axis.minimum_scale=1;chart.value_axis.maximum_scale=100000
    scaling=chart.value_axis._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}scaling")
    logarithm=OxmlElement("c:logBase");logarithm.set("val","10");scaling.insert(0,logarithm)
    chart.plots[0].data_labels.number_format="0.00"
    text(s,.86,6.33,7.3,.29,"平均耗时／毫秒（对数轴）；误差与完整规模对照见报告。",11.5,MUTED)
    ratio=float(rows[1]["elapsed_ms_mean"])/float(rows[2]["elapsed_ms_mean"])
    text(s,9.04,2.70,3.45,1.12,f"{ratio:.2f}×",58,ACCENT,True)
    text(s,9.1,3.85,3.3,1.0,"半朴素相对\n索引朴素加速",21,INK)
    text(s,9.1,5.00,3.3,.85,"候选检查\n84,734 → 8,413",19,MUTED)
    note(s,4)

    s=standard(prs,5,"混合策略的价值，是看清取舍")
    text(s,.70,1.76,11.98,.46,"图级错误闭环20%｜同一自适应评分模型｜三个种子均值",18,MUTED)
    models=["hybrid_adaptive_unlocked","hybrid_adaptive_locked","hybrid_adaptive_calibrated"]
    rows=[next(r for r in D["graph"] if r["scenario"]=="false_loop_020" and r["model"]==m) for m in models]
    data=CategoryChartData();data.categories=["不锁定","硬锁定","校准复核"]
    data.add_series("误报率",[float(r["false_positive_rate_mean"]) for r in rows])
    chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.75),Inches(2.43),Inches(7.4),Inches(3.63),data).chart
    setup_chart(chart);chart.value_axis.minimum_scale=0;chart.value_axis.maximum_scale=.6;chart.value_axis.major_unit=.2
    chart.value_axis.tick_labels.number_format="0%";chart.plots[0].data_labels.number_format="0.00%"
    for point,color in zip(chart.series[0].points,[BLUE,GRAY,ACCENT]):point.format.fill.fore_color.rgb=rgb(color)
    text(s,8.85,2.76,3.7,1.08,"−8.66",54,ACCENT,True)
    text(s,8.92,3.92,3.6,.8,"个百分点\n校准相对硬锁定误报率",18,INK)
    text(s,8.92,4.93,3.6,.8,"F1：0.6144 → 0.6549\n原始分数保持不变",17,MUTED)
    text(s,.86,6.32,11.70,.34,"不锁定F1仍为0.6992；重度扰动下，校准未取得整体优势。",16,INK)
    note(s,5)

    s=standard(prs,6,"结果可复现，结论有边界")
    for x,value,label in [(.84,"200","项目测试通过"),(5.03,"13","项目正式图表"),(9.15,"23,760","图级决策完成回放")]:
        text(s,x,2.0,3.6,.99,value,48,ACCENT,True)
        text(s,x,3.17,3.6,.45,label,18,MUTED)
    connector(s,.84,4.07,12.3,4.07,LINE,1)
    text(s,.86,4.52,11.5,.62,"保留完整证明　·　共用评分入口　·　对照复核策略",23,INK,True)
    text(s,.88,5.43,11.5,.85,"数据与标签均为教学构造。\n下一步：真实资料核验、更多独立重复与跨分布验证。",19,MUTED)
    text(s,11.46,6.37,1.0,.31,"谢谢",16,ACCENT,True,align=PP_ALIGN.RIGHT)
    note(s,6)

    s=standard(prs,7,"理论公式与语义边界",backup=True,kicker="答辩备用 · 理论")
    blocks=[("有限规则闭包","F(t+1) = T_R(F(t))，直到不再产生新事实","有限常量域、安全正Horn规则、无资源截断时讨论完备性。"),
            ("Bayes后验","P(R = 1 | b_O) ∝ π · ∏(j∈O) P(b_j | R = 1)","缺测不按零值进入；星形模型的低开销不能推广到任意网络。"),
            ("混合评分与行动","s = Σ w_i s_i；d_C = 1[s ≥ 0.5] ∨ (A ∧ s ≥ η)","证明始终保留；最低评分支持η只按验证扰动选择。")]
    for i,(label,formula,explanation) in enumerate(blocks):
        y=1.90+i*1.57
        text(s,.85,y,11.75,.38,label,19,ACCENT,True)
        text(s,.88,y+.47,11.5,.49,formula,23,INK)
        text(s,.90,y+1.05,11.5,.35,explanation,15,MUTED)
    note(s,7)

    s=standard(prs,8,"团队责任与复现入口",backup=True,kicker="答辩备用 · 工程与分工")
    table=s.shapes.add_table(6,4,Inches(.72),Inches(1.88),Inches(11.92),Inches(3.62)).table
    for col,w in zip(table.columns,[1.25,2.05,2.28,6.34]):col.width=Inches(w)
    headers=["姓名","学号","职能","责任范围"]
    for j,value in enumerate(headers):table.cell(0,j).text=value
    for i,m in enumerate(B["members"],1):
        for j,value in enumerate([m["name"],m["id"],m["role"],m["responsibility"]]):table.cell(i,j).text=value
    for i,r in enumerate(table.rows):
        for j,c in enumerate(r.cells):
            c.fill.solid();c.fill.fore_color.rgb=rgb("E8DED2" if i==0 else BG)
            c.margin_left=Inches(.09);c.margin_right=Inches(.08);c.margin_top=Inches(.06);c.margin_bottom=Inches(.04)
            c.vertical_anchor=MSO_ANCHOR.MIDDLE
            for p in c.text_frame.paragraphs:
                p.space_after=Pt(0);p.line_spacing=1.08
                for run in p.runs:
                    run.font.name=FONT;run.font.size=Pt(14 if j<3 else 12.5);run.font.bold=i==0;run.font.color.rgb=rgb(INK)
    text(s,.89,5.80,11.6,.34,"网站：根目录“打开网站.cmd” → 实验二",17,INK)
    text(s,.90,6.27,11.6,.30,"python -B -X utf8 main.py lab02 verify --full-replay",16,MUTED,font="Consolas")
    note(s,8)
    prs.core_properties.title=B["title"];prs.core_properties.author="；".join(m["name"] for m in B["members"])
    prs.core_properties.subject="华南农业大学 人工智能导论 "+GROUP
    path=OUT/(GROUP+"_推理实验答辩.pptx");prs.save(path)
    report={"file":path.name,"slides":len(prs.slides),"main_slides":6,"backup_slides":2,
            "native_chart_count":sum(sh.has_chart for sl in prs.slides for sh in sl.shapes),
            "editable_text_shapes":sum(sh.has_text_frame for sl in prs.slides for sh in sl.shapes),
            "presenter":B["presenter"],"leader":"梁镇贤"}
    (WORK/"slides-built.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False))


if __name__=="__main__":main()
