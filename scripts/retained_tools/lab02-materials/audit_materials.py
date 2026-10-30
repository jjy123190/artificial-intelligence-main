"""Validate delivered Office/PDF files, identity fields, data and layout evidence."""
from pathlib import Path
from datetime import datetime,timezone
import hashlib,json,re,sys,zipfile
import xml.etree.ElementTree as ET
import fitz
from docx import Document
from pptx import Presentation

WORK=Path(__file__).resolve().parent
ROOT=WORK.parents[2]
OUT=ROOT/"deliverables/lab02"
sys.path.insert(0,str(ROOT))
from experiments.lab02.eval.artifacts import execution_manifest,fingerprint


def compact(text):return re.sub(r"\s+","",text)


def office_text(path):
    texts=[]
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None,path
        for name in archive.namelist():
            if name.endswith(".xml"):
                try:texts.extend(ET.fromstring(archive.read(name)).itertext())
                except ET.ParseError:raise AssertionError(name)
    return compact("".join(texts))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    brief=json.loads((WORK/"brief.json").read_text(encoding="utf-8"))
    exported=json.loads((WORK/"office-export.json").read_text(encoding="utf-8"))
    built=json.loads((WORK/"documents-built.json").read_text(encoding="utf-8"))
    report_name=next(r["file"] for r in built if r["kind"]=="report")
    report_stat=next(r for r in exported["documents"] if r["file"]==report_name)
    assert 5000 <= report_stat["word_count"] <= 8000,report_stat
    assert 5000 <= report_stat["main_text_word_count"] <= 8000,report_stat
    assert exported["presentations"][0]["slides"]==8
    assert not exported["presentations"][0]["text_height_flags"]
    results=[]
    for path in sorted(OUT.iterdir()):
        if path.suffix not in (".docx",".pptx",".pdf") or not path.name.startswith(f"第{brief['group']}组_"):continue
        item={"file":path.name,"bytes":path.stat().st_size,"sha256":fingerprint(path)}
        if path.suffix in (".docx",".pptx"):
            text=office_text(path)
            assert "刘奕达" not in text and "梁镇闲" not in text,path
            assert "组长：陈宗波" not in text,path
            assert "截止" not in text,path
            assert not re.search(r"\{\{[A-Z_]+\}\}|TODO|TBD|待填写",text),path
            for member in brief["members"]:
                assert member["name"] in text and member["id"] in text,(path,member)
            assert f"第{brief['group']}组" in text and "刘亦达" in text and "梁镇贤" in text
            item["identity_fields_match"]=True
            if path.name==report_name:
                for heading in ["摘要","关键词","引言","相关理论基础","模型建模","算法设计","结果与分析","创新点总结","不足","参考文献"]:
                    assert heading in text,heading
                for number in ["34133.87","362.81","57.75","0.8575","0.7053","43.94","35.28"]:
                    assert number in text,number
                d=Document(path)
                item["equations"]=len(d.element.xpath(".//m:oMath"))
                item["figures"]=len(d.inline_shapes)
                assert item["equations"]>=8 and item["figures"]==4
                assert abs(d.sections[0].left_margin.cm-3.175)<.01
            if path.suffix==".pptx":
                p=Presentation(path)
                assert len(p.slides)==8
                charts=[s.chart for slide in p.slides for s in slide.shapes if s.has_chart]
                assert len(charts)==2
                snapshot=json.loads((WORK/"data-snapshot.json").read_text(encoding="utf-8"))["tables"]
                expected=[float(next(r for r in snapshot["forward"] if r["groups"]=="16" and r["model"]==m)["elapsed_ms_mean"]) for m in ["naive_scan","indexed_naive","semi_naive"]]
                assert all(abs(a-b)<1e-9 for a,b in zip(charts[0].series[0].values,expected))
                expected=[float(next(r for r in snapshot["graph"] if r["scenario"]=="false_loop_020" and r["model"]==m)["false_positive_rate_mean"]) for m in ["hybrid_adaptive_unlocked","hybrid_adaptive_locked","hybrid_adaptive_calibrated"]]
                assert all(abs(a-b)<1e-12 for a,b in zip(charts[1].series[0].values,expected))
                item["native_charts_match_csv"]=True
                item["slides"]=8
        else:
            pdf=fitz.open(path);pages=[]
            all_text=""
            for i,page in enumerate(pdf):
                text=page.get_text();all_text+=text
                assert len(text.strip())>40,(path,i)
                bad=[w for w in page.get_text("words") if w[0]<-1 or w[1]<-1 or w[2]>page.rect.width+1 or w[3]>page.rect.height+1]
                assert not bad,(path,i,bad[:2])
                pages.append({"page":i+1,"text_chars":len(text.strip())})
            assert "刘奕达" not in compact(all_text) and "梁镇闲" not in compact(all_text)
            item["pages"]=len(pdf);item["page_checks"]=pages
            if path.stem==Path(report_name).stem:
                starts=[i for i,page in enumerate(pdf) if "import json" in page.get_text()]
                ends=[i for i,page in enumerate(pdf) if "print(result" in page.get_text()]
                assert starts and ends and starts==ends,(starts,ends)
                item["code_block_kept_on_one_page"]=True
        results.append(item)
    assert len(results)==8,len(results)
    references=json.loads((WORK/"references-verified.json").read_text(encoding="utf-8"))["records"]
    assert len(references)==7 and all(r["status"]=="verified_metadata" for r in references)
    current=execution_manifest(ROOT);sealed=0
    for folder in ("demo","experiments","end_to_end","benchmarks","governance"):
        base=ROOT/"outputs/lab02"/folder
        manifest=json.loads((base/"manifest.json").read_text(encoding="utf-8"))
        assert current["source_sha256"]==manifest["source_sha256"]
        assert current["snapshot_sha256"]==manifest["snapshot_sha256"]
        for name,digest in manifest["artifact_sha256"].items():assert fingerprint(base/name)==digest;sealed+=1
    talk=(WORK/"talk-source.md").read_text(encoding="utf-8")
    segments=re.findall(r"#### 第\d页[^\n]*\n\n(.*?)(?=\n####|\n###|\Z)",talk,re.S)
    assert len(segments)==6
    summary={"completed_utc":datetime.now(timezone.utc).isoformat(),"successful":True,
             "school":brief["school"],"group":int(brief["group"]),"leader":"梁镇贤","presenter":"刘亦达",
             "report_word_count":report_stat["word_count"],"main_text_word_count":report_stat["main_text_word_count"],
             "report_pages":report_stat["pages"],"references_verified":7,"main_slides":6,"backup_slides":2,
             "speech_cjk_characters":sum(len(re.findall(r"[\u4e00-\u9fff]",s)) for s in segments),
             "speech_duration_plan_seconds":165,"actual_rehearsal_not_claimed":True,
             "source_data_and_sealed_results_unchanged":True,"sealed_artifacts":sealed,
             "submission_zip_created":False,"uploaded":False,"files":results}
    (WORK/"materials-audit.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({k:v for k,v in summary.items() if k!="files"},ensure_ascii=False))


if __name__=="__main__":main()
