"""Build the current submission materials from the report's verified release.

All intermediate files stay here; publication copies only reviewed deliverables
to the user's submission directory. The submitted source and report are read-only.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
import shutil
import sys

WORK = Path(__file__).resolve().parent
ROOT = WORK.parents[2]
REPORT_WORK = ROOT / "outputs/qa/lab02-report-rewrite-20261026"
ARTIFACTS = WORK / "artifacts"
PPT_WORK = WORK / "ppt-video"
DEMO_WORK = WORK / "demo-video"
SESSION = "01a177c9-84f6-7003-9cf2-2692d2a573ae"
TABLE_FILES = {
    "forward": "benchmarks/forward_summary.csv",
    "slicing": "benchmarks/goal_slicing.csv",
    "resolution": "benchmarks/resolution.csv",
    "indicator": "experiments/summary.csv",
    "sensitivity": "experiments/sensitivity.csv",
    "ablation": "experiments/ablations.csv",
    "conflict": "experiments/conflict_summary.csv",
    "graph": "end_to_end/summary.csv",
    "equity": "governance/equity_summary.csv",
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def row(evidence, family, **conditions):
    return next(r for r in evidence["tables"][family]
                if all(r[k] == v for k, v in conditions.items()))


def freeze():
    evidence = read_json(REPORT_WORK / "evidence.json")
    release = Path(evidence["release"])
    destination = release.parent
    for folder in (ARTIFACTS, PPT_WORK, DEMO_WORK):
        folder.mkdir(parents=True, exist_ok=True)
    for name, expected in evidence["source_sha256"].items():
        assert digest(release / name) == expected, name
    for name, expected in evidence["raw_data_sha256"].items():
        assert digest(release / "data/raw" / name) == expected, name
    for family, name in TABLE_FILES.items():
        with (release / "outputs/lab02" / name).open(encoding="utf-8-sig", newline="") as handle:
            assert list(csv.DictReader(handle)) == evidence["tables"][family], name
    manifest_files = {}
    for family in ("demo", "experiments", "end_to_end", "benchmarks", "governance"):
        base = release / "outputs/lab02" / family
        manifest = read_json(base / "manifest.json")
        assert manifest["source_sha256"] == evidence["source_sha256"], family
        assert manifest["snapshot_sha256"] == evidence["raw_data_sha256"], family
        for name, expected in manifest["artifact_sha256"].items():
            assert digest(base / name) == expected, (family, name)
        manifest_files[family] = digest(base / "manifest.json")
    report = read_json(REPORT_WORK / "completion.json")
    for asset in report["published"]:
        assert digest(asset["path"]) == asset["sha256"], asset["path"]
    assert evidence["tests"]["tests"] == 204 and evidence["tests"]["passed"]
    assert evidence["acceptance"]["successful"]
    write_json(WORK / "evidence.json", evidence)
    snapshot = {
        "source_session": SESSION, "release": str(release),
        "source_files": len(evidence["source_sha256"]), "data_files": len(evidence["raw_data_sha256"]),
        "tables": len(TABLE_FILES), "project_tests": 204,
        "manifest_sha256": manifest_files,
        "protected_reports": {a["path"]: a["sha256"] for a in report["published"]},
        "destination_before": {p.name: digest(p) for p in destination.iterdir() if p.is_file()},
        "source_unchanged": True, "tables_match_report": True,
        "requirements": {"main_slides": 9, "supplementary_slides": 4,
                         "ppt_duration_seconds": 180, "demo_duration_seconds": 60,
                         "ppt_audio": True, "ppt_subtitles": True,
                         "demo_audio": False, "demo_subtitles": False,
                         "visible_ppt_source_footers": False},
    }
    write_json(WORK / "source-check.json", snapshot)
    print(json.dumps({k: snapshot[k] for k in ["source_files", "data_files", "tables", "project_tests", "tables_match_report"]}), flush=True)


def narration():
    evidence = read_json(WORK / "evidence.json")
    old = ROOT / "outputs/qa/lab02-ppt-video-3min"
    settings = read_json(old / "narration.json")
    indexed = float(row(evidence, "forward", groups="16", model="indexed_naive")["elapsed_ms_mean"])
    semi = float(row(evidence, "forward", groups="16", model="semi_naive")["elapsed_ms_mean"])
    # The displayed values are generated from CSV; their spoken forms are checked
    # explicitly so the synthesized narration cannot silently retain the old run.
    assert f"{indexed:.2f}" == "1062.37" and f"{semi:.2f}" == "181.41"
    assert f"{indexed / semi:.2f}" == "5.86"
    settings["slides"][4]["cues"][2] = {
        "text": f"索引朴素到半朴素，耗时从{indexed:.2f}降到{semi:.2f}毫秒，",
        "speech": "索引朴素到半朴素，耗时从一千零六十二点三七降到一百八十一点四一毫秒，",
    }
    settings["slides"][4]["cues"][3] = {
        "text": f"加速约{indexed / semi:.2f}倍；迭代轮数不变，候选检查显著减少。",
        "speech": "加速约五点八六倍；迭代轮数不变，候选检查显著减少。",
    }
    settings["slides"][8]["cues"][2] = {
        "text": "项目204项测试通过，23760条图级决策已回放。",
        "speech": "项目二百零四项测试通过，两万三千七百六十条图级决策已回放。",
    }
    write_json(PPT_WORK / "narration.json", settings)
    for p in old.glob("voice-*"):
        if p.suffix in (".mp3", ".json"):
            target = PPT_WORK / p.name
            if not target.exists():
                shutil.copy2(p, target)
    print("Narration updated; unchanged speech is reusable through content hashes.", flush=True)


def innovation():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    e = read_json(WORK / "evidence.json")
    brief = read_json(REPORT_WORK / "brief.json")
    d = Document()
    section = d.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.05)
    section.left_margin = section.right_margin = Cm(2.15)
    section.header_distance = section.footer_distance = Cm(.9)

    def font(run, size=11.5, bold=False, color="292724", east="宋体"):
        run.font.name = "Times New Roman"
        run.font.size, run.font.bold = Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east)

    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(11.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.30
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name in ["Heading 1", "Heading 2"]:
        style = d.styles[name]
        style.font.size, style.font.bold = Pt(14 if name == "Heading 1" else 12), True
        style.font.color.rgb = RGBColor.from_string("9A4728")
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(11), Pt(6)
        style.paragraph_format.keep_with_next = True

    def para(value, size=11.5, indent=True, color="292724"):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Pt(23 if indent else 0)
        font(p.add_run(value), size, color=color)
        return p

    def title(value, size, color="292724"):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(9)
        font(p.add_run(value), size, True, color, "黑体")

    def table(headers, rows, widths):
        t = d.add_table(rows=1, cols=len(headers))
        t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
        for values in [headers, *rows]:
            r = t.rows[0] if values is headers else t.add_row()
            r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if values is headers:
                r._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for cell, value, width in zip(r.cells, values, widths):
                cell.width, cell.vertical_alignment = Cm(width), WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.12
                font(p.add_run(str(value)), 10.5, values is headers)
                if values is headers:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EEE8E0")
                    cell._tc.get_or_add_tcPr().append(shade)
        d.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(header.add_run("华南农业大学｜人工智能导论｜软件工程5班·第31组"), 9, color="71685F")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    title("创新点说明", 22)
    title("面向企业股权风险评估的确定性与不确定性混合推理", 13)
    p = para("第31组　梁镇贤、陈宗波、温文辉、陈宇冲、刘亦达\n2026年10月26日", 10.5, False, "71685F")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_heading("一、问题与创新定位", 1)
    para("股权、交易和担保关系可以形成明确的结构证明，但输入关系可能缺失或错误；概率与模糊评分能够吸收不确定指标，却不能代替规则证明。若将两类结果直接合成一个风险分数，就难以辨别结论来自证据、模型还是人为锁定。")
    para("本项目选择课程的“确定性＋不确定性混合推理”方向，将结构证明、数值评分和最终复核行动分开保存，在同一教学案例上比较不锁定、结构硬锁定与验证集校准三种策略。创新集中于可审计的融合流程、行动校准和公平对照；半朴素求值、Bayes与Mamdani属于已有方法的复现和适配，不宣称首创这些算法。")
    d.add_heading("二、核心设计：保留证明，校准行动", 1)
    para("令 A 表示由观测关系图推出的高复核预警，s∈[0,1] 表示共同评分模型的输出，η为额外结构行动所需的最低评分支持。三种策略使用相同输入事实、相同原始评分和0.5分类阈值，仅改变最终行动：")
    table(["策略", "复核行动", "可追溯信息"], [
        ["不锁定", "dU = 1[s ≥ 0.5]", "原始评分与结构事实"],
        ["结构硬锁定", "dL = dU ∨ A", "强制保留结构预警"],
        ["校准复核", "dC = dU ∨ (A ∧ s ≥ η)", "保留证明，允许暂缓额外行动"],
    ], [2.8, 6.2, 7.6])
    para("校准策略从{0、0.15、0.25、0.35、0.45、0.5}选择η，在预设验证扰动上最小化(FP＋2FN)/N；代价相同则选择较低η。每个种子的320条校准记录均来自验证阶段，测试标签不参与选择。三个种子7、19、43分别选出0.25、0.15、0，说明校准可能退化为硬锁定，不能假定每次都改善结果。")
    para("实现将score与review_required分开，并保存provided_facts、retained_proofs和策略选择记录。暂缓复核不会将已证线索改称“未证明”，结构相容性也不等同于真实风险准确率。", 10.5)

    d.add_page_break()
    d.add_heading("三、量化验证：误报、召回与效率分别评价", 1)
    para("图级实验以完整关系图为输入，覆盖3个种子、6种观测条件和11种方法，共保存23,760条决策。下表使用20%错误闭环条件下的同一自适应评分模型；各单元为三个种子的均值±样本标准差。")
    models = ["hybrid_adaptive_unlocked", "hybrid_adaptive_locked", "hybrid_adaptive_calibrated"]
    graph_rows = [row(e, "graph", scenario="false_loop_020", model=model) for model in models]
    metrics = []
    for label, r in zip(["不锁定", "硬锁定", "校准复核"], graph_rows):
        metrics.append([label,
            f"{100*float(r['false_positive_rate_mean']):.2f}±{100*float(r['false_positive_rate_sd']):.2f}",
            f"{100*float(r['recall_mean']):.2f}±{100*float(r['recall_sd']):.2f}",
            f"{float(r['f1_mean']):.4f}±{float(r['f1_sd']):.4f}"])
    table(["行动策略", "误报率／%", "召回率／%", "F1"], metrics, [2.8, 4.3, 4.3, 5.2])
    unlocked, locked, calibrated = graph_rows
    drop = 100 * (float(locked["false_positive_rate_mean"]) - float(calibrated["false_positive_rate_mean"]))
    para(f"校准相对硬锁定的误报率由43.94%降至35.28%，减少{drop:.2f}个百分点；F1由0.6144升至0.6549，召回率则由87.71%降至86.16%。按相同漏报代价计算的平均决策代价由0.3750降至0.3306，但仍高于不锁定的0.2694。结果体现了复核负担与漏报之间的取舍。")
    para("校准仍未全面超过不锁定：其误报率更高，F1更低；真实结构预警保留率也从硬锁定的100%降至94.41%。三个种子的波动较大，本实验只报告描述性结果，不作统计显著性或真实企业有效性的宣称。")
    d.add_heading("确定性优化：同闭包条件下比较计算开销", 2)
    para("基准覆盖1、4、8、16组结构，每种方法重复3次。16组时三种算法均得到3,031条事实、20轮迭代；先核对闭包一致，再比较候选检查与耗时。下表为最新提交源码对应的同批实测。")
    bench = [row(e, "forward", groups="16", model=model) for model in ["naive_scan", "indexed_naive", "semi_naive"]]
    table(["方法", "平均耗时±标准差／ms", "候选检查次数"], [
        [label, f"{float(r['elapsed_ms_mean']):,.2f}±{float(r['elapsed_ms_sd']):,.2f}", f"{int(float(r['candidate_checks_mean'])):,}"]
        for label, r in zip(["全扫描", "索引朴素", "半朴素"], bench)
    ], [3.2, 7.0, 6.4])
    speedup = float(bench[1]["elapsed_ms_mean"]) / float(bench[2]["elapsed_ms_mean"])
    para(f"半朴素通过增量事实驱动减少重复连接，相对索引朴素耗时从1,062.37 ms降至181.41 ms，加速约{speedup:.2f}倍；候选检查从84,734降至8,413。加速来自计算量减少，未改变迭代轮数或闭包语义；具体耗时受硬件和负载影响。")

    d.add_page_break()
    d.add_heading("四、不确定性优化与消融解释", 1)
    para("两种主方法均使用同一企业指标。Bayes从训练集估计先验与CPT；Mamdani在训练集拟合15组候选参数，缺测时向中性值收缩。权重和策略在测试前冻结。指标实验覆盖3个种子、8种条件，每种子使用1,000条训练、300条验证和600条测试样本。")
    para("清洁条件下，学习Bayes的F1由手工参数的0.7767提高至0.8575；拟合Mamdani的F1为0.7053。最新单样本耗时均值分别为6.7 μs和245.8 μs。严重噪声σ=0.35且缺测40%时，学习Bayes的F1降至0.6174，参数学习并未消除分布扰动的影响。")
    para("18条消融操作中，12条为移除零权重组件后不变，6条为显式回退基线。清洁指标实验的调优权重退化为纯Bayes，因此不把相同F1解释为融合增益，也不把回退后的差值冒充组件边际贡献。D-S证据冲突折扣和循环权益求解是额外扩展，未替代两种主方法或本项混合创新。")
    d.add_heading("五、工程实现与复现依据", 1)
    para("最新提交代码兼容v2、v3训练模型的命令行和网页重载，补充了算法及参数选择注释；204项项目测试和8组独立数学参照检查通过。完整保存结果验收已重放201,600次指标评分、43,200次D-S评分和23,760条图级决策；本说明与新版报告使用同一源码及结果快照。")
    table(["内容", "源码／结果位置（相对源码目录）"], [
        ["结构求值与优化", "experiments/lab02/deterministic/forward.py\noutputs/lab02/benchmarks/forward_summary.csv"],
        ["评分、决策与校准", "experiments/lab02/hybrid/model.py\nexperiments/lab02/hybrid/decision.py\nexperiments/lab02/hybrid/calibration.py"],
        ["图级策略验证", "outputs/lab02/end_to_end/summary.csv\noutputs/lab02/end_to_end/models.json"],
        ["参数及消融", "outputs/lab02/experiments/summary.csv\noutputs/lab02/experiments/ablations.csv"],
    ], [3.4, 13.2])
    para("在源码目录中，双击“打开网站.cmd”或“启动实验二.cmd”进入网站；“运行实验.cmd”菜单可重新生成完整结果。命令行使用python main.py lab02 all生成实验，再以python main.py lab02 verify --full-replay核对结果。环境配置与结果解释见新版实验报告附录B，首次生成结果和安装依赖需要等待。", 10.5)
    d.add_heading("六、适用范围与后续工作", 1)
    para("全部实体、输入关系和高复核需求标签均为教学构造，控制候选、经济权益、模型评分与复核行动各有语义。当前贡献是课程场景中的可解释实现与可验证对照，不能直接用于法律控制认定或真实企业风险结论。后续应在有授权且可核验的数据上测试跨分布稳定性、结构事实可靠性及实际人工复核成本。")
    d.core_properties.title = "第31组_创新点说明"
    d.core_properties.subject = brief["title"]
    d.core_properties.author = "；".join(m["name"] for m in brief["members"])
    d.core_properties.keywords = "确定性推理,不确定性推理,可审计混合推理,校准复核"
    destination = ARTIFACTS / "第31组_创新点说明.docx"
    d.save(destination)
    text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    (WORK / "innovation-text.txt").write_text(text, encoding="utf-8")
    assert "6.28" not in text and "200项" not in text
    write_json(WORK / "innovation-build.json", {"file": destination.name, "tables": len(d.tables), "characters": len(text), "sha256": digest(destination)})
    print("Built", destination, flush=True)


def export_innovation():
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    app = doc = None
    path = ARTIFACTS / "第31组_创新点说明.docx"
    before = digest(path)
    try:
        app = win32com.client.DispatchEx("Word.Application")
        app.Visible = False
        app.DisplayAlerts = 0
        doc = app.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False)
        doc.Repaginate()
        pages = doc.ComputeStatistics(2)
        words = doc.ComputeStatistics(0)
        doc.ExportAsFixedFormat(str(path.with_suffix(".pdf")), 17)
        write_json(WORK / "innovation-export.json", {"pages": pages, "words": words, "source_unchanged": digest(path) == before})
        print("Innovation PDF exported:", pages, "pages,", words, "words", flush=True)
    finally:
        if doc is not None:
            doc.Close(False)
        if app is not None:
            try:
                if app.Documents.Count == 0:
                    app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("stage", choices=["freeze", "narration", "innovation", "export-innovation"])
    args = parser.parse_args()
    {"freeze": freeze, "narration": narration, "innovation": innovation,
     "export-innovation": export_innovation}[args.stage]()


if __name__ == "__main__":
    main()
